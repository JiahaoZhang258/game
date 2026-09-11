"""Question 2 rebuilt solution.

The policy is causal in two layers:

* At 00:00, the plan uses only completed days.
* During the day, the fixed purchase plan is executed one 10-minute interval
  at a time with the current observation and the planned SOC reference.

January is simulated physically from the stated 6000 kWh initial SOC. Its
records are not reported, but its ending SOC is passed to February 1.
"""
from __future__ import annotations

import argparse
import json
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import coo_matrix

from solve_q1 import (
    DATA,
    DT,
    ETA,
    P_MAX,
    ROOT,
    SOC_MAX,
    SOC_MIN,
    T,
    four_hour_bins,
    load_attachment1,
)


TEMPLATE = DATA / "附件5" / "result2.xlsx"
INITIAL_SOC = 6000.0
JANUARY_START = date(2025, 1, 1)
JANUARY_END = date(2025, 1, 31)
JANUARY_VALIDATION_START = date(2025, 1, 22)
REPORT_START = date(2025, 2, 1)
REPORT_END = date(2025, 12, 31)
SELECTED_DATES = [
    date(2025, 3, 20),
    date(2025, 6, 21),
    date(2025, 9, 23),
    date(2025, 12, 21),
]

ALPHA = 0.825
ERROR_WINDOW = 7
RESERVE_RATIO = 0.625
TERMINAL_VALUE = 0.30
TERMINAL_SOC_TARGET = 2400.0
REGULARIZER = 1e-5
EMERGENCY_EPS = 1e-7


@dataclass(frozen=True)
class Parameters:
    alpha: float = ALPHA
    error_window: int = ERROR_WINDOW
    reserve_ratio: float = RESERVE_RATIO
    terminal_value: float = TERMINAL_VALUE
    terminal_soc_target: float | None = TERMINAL_SOC_TARGET


WARMUP_PARAMETERS = Parameters(
    alpha=0.80, error_window=7, reserve_ratio=0.0,
    terminal_value=0.60, terminal_soc_target=None,
)


PARAMETER_CANDIDATES = tuple(
    Parameters(alpha=alpha, error_window=window, reserve_ratio=reserve,
               terminal_value=terminal, terminal_soc_target=None)
    for alpha in (0.75, 0.80)
    for window in (4, 7)
    for reserve in (0.0, 0.05)
    for terminal in (0.30, 0.60)
)


def _as_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date(1899, 12, 30) + timedelta(days=int(float(value)))


def load_attachment2() -> tuple[list[date], np.ndarray, np.ndarray]:
    wb = load_workbook(DATA / "附件2.xlsx", read_only=True, data_only=True)

    def read_sheet(name: str) -> tuple[list[date], np.ndarray]:
        ws = wb[name]
        rows = ws.iter_rows(values_only=True)
        next(rows)
        dates, values = [], []
        for row in rows:
            if row[0] is None:
                continue
            dates.append(_as_date(row[0]))
            values.append([float(x or 0.0) for x in row[1 : T + 1]])
        return dates, np.asarray(values, dtype=float)

    load_dates, load = read_sheet("小区负载")
    pv_dates, pv = read_sheet("光伏发电实际功率")
    wb.close()
    if load_dates != pv_dates or load.shape != pv.shape or load.shape[1] != T:
        raise ValueError("附件2的负荷和光伏日期或维度不一致")
    return load_dates, load, pv


def _weighted_mean(rows: list[np.ndarray]) -> np.ndarray:
    if not rows:
        raise ValueError("at least one row is required")
    weights = np.arange(1, len(rows) + 1, dtype=float)
    weights /= weights.sum()
    return np.average(np.asarray(rows), axis=0, weights=weights)


def _smooth(values: np.ndarray) -> np.ndarray:
    padded = np.pad(values, (1, 1), mode="edge")
    return np.convolve(padded, np.ones(3) / 3.0, mode="valid")


def load_history_forecast(day_index: int, dates: list[date], load: np.ndarray, typical_load: np.ndarray) -> np.ndarray:
    """Forecast load from the last four completed same-weekday rows."""
    if day_index == 0:
        return _smooth(typical_load.copy())
    history = [
        load[i]
        for i in range(day_index)
        if dates[i].weekday() == dates[day_index].weekday()
    ][-4:]
    source = np.mean(history, axis=0) if history else np.mean(load[:day_index], axis=0)
    return _smooth(source)


def pv_trend_forecast(day_index: int, pv: np.ndarray, typical_pv: np.ndarray) -> np.ndarray:
    """Causal 21-day linear-trend PV forecast."""
    if day_index == 0:
        return np.maximum(typical_pv.copy(), 0.0)
    history = pv[max(0, day_index - 21):day_index]
    if len(history) == 1:
        return np.maximum(history[0].copy(), 0.0)
    x = np.arange(len(history), dtype=float)
    centered = x - x.mean()
    slope = centered @ np.asarray(history, dtype=float) / (centered @ centered)
    return np.maximum(np.mean(history, axis=0) + slope * (len(history) - x.mean()), 0.0)


def forecast_day(
    day_index: int,
    dates: list[date],
    load: np.ndarray,
    pv: np.ndarray,
    typical_load: np.ndarray,
    typical_pv: np.ndarray,
    parameters: Parameters,
    forecast_variant: str = "optimized",
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Return forecast load, forecast PV, and causal net-demand residuals."""
    def basic_forecast(j: int) -> tuple[np.ndarray, np.ndarray]:
        if forecast_variant == "official":
            if j == 0:
                return typical_load.copy(), typical_pv.copy()
            weekday_history = [load[i] for i in range(j) if dates[i].weekday() == dates[j].weekday()][-4:]
            load_hat = _smooth(_weighted_mean(weekday_history) if weekday_history else np.mean(load[:j], axis=0))
            recent_pv = list(pv[max(0, j - 4):j])
            pv_hat = _weighted_mean(recent_pv) if recent_pv else typical_pv.copy()
            return load_hat, np.maximum(pv_hat, 0.0)
        if forecast_variant != "optimized":
            raise ValueError(f"unknown forecast_variant: {forecast_variant}")
        return load_history_forecast(j, dates, load, typical_load), pv_trend_forecast(j, pv, typical_pv)

    load_f, pv_f = basic_forecast(day_index)
    if day_index == 0:
        return load_f, pv_f, np.empty((0, T))
    residuals = []
    start = max(1, day_index - parameters.error_window)
    for j in range(start, day_index):
        previous_load, previous_pv = basic_forecast(j)
        actual_net = (load[j] - pv[j]) * DT
        predicted_net = (previous_load - previous_pv) * DT
        residuals.append(actual_net - predicted_net)
    return load_f, pv_f, np.asarray(residuals, dtype=float)


def design_demand(
    load_f: np.ndarray,
    pv_f: np.ndarray,
    residuals: np.ndarray,
    alpha: float,
    recent_decay: float | None = None,
) -> np.ndarray:
    predicted = (load_f - pv_f) * DT
    if residuals.size == 0:
        return np.maximum(predicted, 0.0)
    if recent_decay is None:
        correction = np.quantile(residuals, alpha, axis=0, method="linear")
    else:
        if not 0.0 < recent_decay <= 1.0:
            raise ValueError("recent_decay must be in (0, 1]")
        weights = recent_decay ** np.arange(len(residuals) - 1, -1, -1, dtype=float)
        order = np.argsort(residuals, axis=0)
        sorted_values = np.take_along_axis(residuals, order, axis=0)
        sorted_weights = np.take_along_axis(np.broadcast_to(weights[:, None], residuals.shape), order, axis=0)
        cumulative = np.cumsum(sorted_weights, axis=0)
        threshold = alpha * cumulative[-1]
        correction = np.empty(residuals.shape[1], dtype=float)
        for column in range(residuals.shape[1]):
            correction[column] = sorted_values[np.searchsorted(cumulative[:, column], threshold[column]), column]
    # A negative value means photovoltaic surplus available for charging.
    # Clipping it to zero would make the day-ahead plan buy unnecessary power.
    return predicted + correction


def plan_day(
    price: np.ndarray,
    demand: np.ndarray,
    initial_soc: float,
    parameters: Parameters,
    terminal_soc_target: float | None = None,
) -> dict[str, np.ndarray]:
    """Solve a deterministic planning LP for the risk-adjusted demand."""
    if terminal_soc_target is None:
        terminal_soc_target = parameters.terminal_soc_target
    # q, charge, discharge, soc_end, surplus
    n = 5 * T
    q, charge, discharge, soc, surplus = 0, T, 2 * T, 3 * T, 4 * T
    objective = np.zeros(n)
    objective[q : q + T] = price
    objective[charge : charge + T] = REGULARIZER
    objective[discharge : discharge + T] = REGULARIZER
    objective[soc + T - 1] = -parameters.terminal_value

    bounds = [(0.0, None)] * T
    bounds += [(0.0, P_MAX * DT)] * T
    bounds += [(0.0, P_MAX * DT)] * T
    bounds += [(SOC_MIN, SOC_MAX)] * T
    bounds += [(0.0, None)] * T

    rows, cols, values, rhs = [], [], [], []

    def add(row: int, col: int, value: float) -> None:
        rows.append(row)
        cols.append(col)
        values.append(value)

    row = 0
    for t in range(T):
        add(row, q + t, 1.0)
        add(row, discharge + t, 1.0)
        add(row, charge + t, -1.0)
        add(row, surplus + t, -1.0)
        rhs.append(float(demand[t]))
        row += 1

    if terminal_soc_target is not None:
        add(row, soc + T - 1, 1.0)
        rhs.append(float(terminal_soc_target))
        row += 1

    for t in range(T):
        add(row, soc + t, 1.0)
        add(row, charge + t, -ETA)
        add(row, discharge + t, 1.0 / ETA)
        if t:
            add(row, soc + t - 1, -1.0)
            rhs.append(0.0)
        else:
            rhs.append(float(initial_soc))
        row += 1

    result = linprog(
        objective,
        A_eq=coo_matrix((values, (rows, cols)), shape=(row, n)).tocsr(),
        b_eq=np.asarray(rhs),
        bounds=bounds,
        method="highs",
    )
    if not result.success:
        raise RuntimeError(f"日前计划 LP 求解失败: {result.message}")
    x = result.x
    return {
        "purchase_kw": np.maximum(x[q : q + T], 0.0),
        "charge_kw": np.maximum(x[charge : charge + T], 0.0),
        "discharge_kw": np.maximum(x[discharge : discharge + T], 0.0),
        "soc": x[soc : soc + T],
        "surplus_kw": np.maximum(x[surplus : surplus + T], 0.0),
    }


def execute_day(
    price: np.ndarray,
    actual_load: np.ndarray,
    actual_pv: np.ndarray,
    purchase_kw: np.ndarray,
    reference_soc: np.ndarray,
    initial_soc: float,
    reserve_ratio: float,
) -> dict[str, np.ndarray]:
    """Causal interval feedback. The purchase plan remains fixed all day."""
    charge = np.zeros(T)
    discharge = np.zeros(T)
    emergency = np.zeros(T)
    surplus = np.zeros(T)
    soc_trace = np.zeros(T)
    state = float(initial_soc)

    for t in range(T):
        net = (float(actual_load[t]) - float(actual_pv[t])) * DT
        planned_energy = float(purchase_kw[t])
        balance = planned_energy - net
        if balance >= 0.0:
            charge_energy = min(
                balance,
                P_MAX * DT,
                max((SOC_MAX - state) / ETA, 0.0),
            )
            charge[t] = charge_energy
            surplus[t] = balance - charge_energy
        else:
            reserve = SOC_MIN + reserve_ratio * max(
                float(reference_soc[t]) - SOC_MIN, 0.0
            )
            discharge_energy = min(
                -balance,
                P_MAX * DT,
                max(ETA * (state - reserve), 0.0),
            )
            discharge[t] = discharge_energy
            emergency[t] = -balance - discharge_energy
        state += ETA * charge[t] - discharge[t] / ETA
        state = float(np.clip(state, SOC_MIN, SOC_MAX))
        soc_trace[t] = state

    return {
        "charge_kwh": charge,
        "discharge_kwh": discharge,
        "emergency_kwh": emergency,
        "surplus_kwh": surplus,
        "soc_trace": soc_trace,
    }


def emergency_segments(emergency_kwh: np.ndarray) -> list[tuple[str, float]]:
    events = []
    start = None
    for t in range(T + 1):
        active = t < T and emergency_kwh[t] > EMERGENCY_EPS
        if active and start is None:
            start = t
        if not active and start is not None:
            events.append((interval_label(start, t), float(emergency_kwh[start:t].sum())))
            start = None
    return events


def source_four_hour_bins() -> list[np.ndarray]:
    """Group source endpoints into six consecutive natural-day blocks."""
    return four_hour_bins((np.arange(T, dtype=int) + 1) * 10)


def interval_label(start_slot: int, end_slot: int) -> str:
    def fmt(slot: int) -> str:
        minutes = slot * 10
        hour, minute = divmod(minutes, 60)
        return f"{hour}:{minute:02d}"

    return f"{fmt(start_slot)}-{fmt(end_slot)}"


def simulate(
    parameters: Parameters,
    report_start: date = REPORT_START,
    report_end: date = REPORT_END,
    warmup_parameters: Parameters | None = None,
    apply_error_correction: bool = True,
    warmup_apply_error_correction: bool | None = None,
    clip_negative_demand: bool = True,
    recent_decay: float | None = None,
    terminal_soc_target: float | None = None,
    forecast_variant: str = "optimized",
    warmup_forecast_variant: str = "official",
) -> tuple[dict, dict]:
    dates, load, pv = load_attachment2()
    att1 = load_attachment1(DATA / "附件1.xlsx")
    price = np.asarray(att1["price"], dtype=float)
    typical_load = np.asarray(att1["load"], dtype=float)
    typical_pv = np.asarray(att1["pv"], dtype=float)

    if report_start < dates[0] or report_end > dates[-1] or report_start > report_end:
        raise ValueError("模拟日期范围超出附件2或起止日期无效")

    records = []
    state = INITIAL_SOC
    state_at_report_start = None
    for i, day in enumerate(dates):
        if day > report_end:
            break
        active_parameters = (warmup_parameters or WARMUP_PARAMETERS) if day < report_start else parameters
        active_forecast_variant = (
            warmup_forecast_variant if day < report_start else forecast_variant
        )
        load_f, pv_f, residuals = forecast_day(
            i, dates, load, pv, typical_load, typical_pv, active_parameters,
            forecast_variant=active_forecast_variant,
        )
        use_error_correction = (
            warmup_apply_error_correction
            if day < report_start and warmup_apply_error_correction is not None
            else apply_error_correction
        )
        if use_error_correction:
            active_decay = None if day < report_start else recent_decay
            demand = design_demand(load_f, pv_f, residuals, active_parameters.alpha, active_decay)
        else:
            raw_demand = (load_f - pv_f) * DT
            demand = np.maximum(raw_demand, 0.0) if clip_negative_demand else raw_demand
        target = terminal_soc_target if day >= report_start else active_parameters.terminal_soc_target
        plan = plan_day(price, demand, state, active_parameters, target)
        execution = execute_day(
            price,
            load[i],
            pv[i],
            plan["purchase_kw"],
            plan["soc"],
            state,
            active_parameters.reserve_ratio,
        )
        record = {
            "date": day,
            "soc0": float(state),
            "soc24": float(execution["soc_trace"][-1]),
            "purchase_kwh": plan["purchase_kw"],
            "charge_kwh": execution["charge_kwh"],
            "discharge_kwh": execution["discharge_kwh"],
            "emergency_kwh": execution["emergency_kwh"],
            "surplus_kwh": execution["surplus_kwh"],
            "soc_trace": execution["soc_trace"],
            "load_kwh": load[i] * DT,
            "pv_kwh": pv[i] * DT,
            "design_demand_kwh": demand,
            "reference_soc": plan["soc"],
        }
        if day == report_start:
            state_at_report_start = float(state)
        if report_start <= day <= report_end:
            records.append(record)
        state = float(execution["soc_trace"][-1])

    if state_at_report_start is None:
        raise RuntimeError(f"没有找到模拟起点 {report_start}")
    bundle = {"price": price, "records": records}
    summary = summarize(bundle)
    summary["state_at_report_start"] = state_at_report_start
    summary["period_start"] = report_start.isoformat()
    summary["period_end"] = report_end.isoformat()
    summary["period_end_soc"] = float(records[-1]["soc24"])
    return bundle, summary


def select_parameters_from_january() -> tuple[Parameters, dict]:
    """Select and freeze parameters using only the last ten days of January."""
    results = []
    for candidate in PARAMETER_CANDIDATES:
        _, summary = simulate(
            candidate,
            report_start=JANUARY_VALIDATION_START,
            report_end=JANUARY_END,
            warmup_parameters=WARMUP_PARAMETERS,
            forecast_variant="official",
        )
        results.append(
            {
                "parameters": asdict(candidate),
                "total_cost": summary["total_cost"],
                "emergency_kwh": summary["emergency_kwh"],
                "emergency_days": summary["emergency_days"],
            }
        )
    selected_row = min(
        results,
        key=lambda row: (row["total_cost"], row["emergency_kwh"], row["emergency_days"]),
    )
    selected = Parameters(**selected_row["parameters"])
    return selected, {
        "method": "January walk-forward validation",
        "training_period": f"{JANUARY_START.isoformat()} to {JANUARY_VALIDATION_START - timedelta(days=1)}",
        "validation_period": f"{JANUARY_VALIDATION_START.isoformat()} to {JANUARY_END.isoformat()}",
        "candidate_count": len(results),
        "selected": asdict(selected),
        "candidates": results,
    }


def causal_smoke_test(parameters: Parameters) -> dict:
    """Check that future observations cannot change earlier decisions."""
    dates, load, pv = load_attachment2()
    att1 = load_attachment1(DATA / "附件1.xlsx")
    price = np.asarray(att1["price"], dtype=float)
    typical_load = np.asarray(att1["load"], dtype=float)
    typical_pv = np.asarray(att1["pv"], dtype=float)
    target = dates.index(REPORT_START)

    base_load_f, base_pv_f, base_residuals = forecast_day(
        target, dates, load, pv, typical_load, typical_pv, parameters,
        forecast_variant="optimized",
    )
    altered_load = load.copy()
    altered_pv = pv.copy()
    altered_load[target:] += 1234.5
    altered_pv[target:] = np.maximum(altered_pv[target:] - 432.1, 0.0)
    future_load_f, future_pv_f, future_residuals = forecast_day(
        target, dates, altered_load, altered_pv, typical_load, typical_pv, parameters
    )
    base_demand = design_demand(base_load_f, base_pv_f, base_residuals, parameters.alpha)
    future_demand = design_demand(future_load_f, future_pv_f, future_residuals, parameters.alpha)
    if not (
        np.array_equal(base_load_f, future_load_f)
        and np.array_equal(base_pv_f, future_pv_f)
        and np.array_equal(base_demand, future_demand)
    ):
        raise AssertionError("目标日前后的数据改变了 2 月 1 日日前计划输入")

    plan = plan_day(price, base_demand, INITIAL_SOC, parameters, TERMINAL_SOC_TARGET)
    prefix = 36
    base_exec = execute_day(
        price, load[target], pv[target], plan["purchase_kw"], plan["soc"], INITIAL_SOC, parameters.reserve_ratio
    )
    altered_day_load = load[target].copy()
    altered_day_pv = pv[target].copy()
    altered_day_load[prefix:] += 777.0
    altered_day_pv[prefix:] = np.maximum(altered_day_pv[prefix:] - 222.0, 0.0)
    altered_exec = execute_day(
        price, altered_day_load, altered_day_pv, plan["purchase_kw"], plan["soc"], INITIAL_SOC, parameters.reserve_ratio
    )
    for key in ("charge_kwh", "discharge_kwh", "emergency_kwh", "surplus_kwh", "soc_trace"):
        if not np.array_equal(base_exec[key][:prefix], altered_exec[key][:prefix]):
            raise AssertionError(f"当天未来实测值改变了已执行前缀: {key}")
    return {
        "status": "PASS",
        "checks": [
            "future dates do not change February 1 day-ahead inputs",
            "future intraday observations do not change executed prefix",
        ],
    }


def summarize(bundle: dict) -> dict:
    price = bundle["price"]
    records = bundle["records"]
    plan_kwh = sum(float(r["purchase_kwh"].sum()) for r in records)
    emergency_kwh = sum(float(r["emergency_kwh"].sum()) for r in records)
    plan_cost = sum(float(np.dot(price, r["purchase_kwh"])) for r in records)
    emergency_cost = sum(float(np.dot(5.0 * price, r["emergency_kwh"])) for r in records)
    return {
        "days": len(records),
        "intervals": len(records) * T,
        "plan_kwh": plan_kwh,
        "emergency_kwh": emergency_kwh,
        "plan_cost": plan_cost,
        "emergency_cost": emergency_cost,
        "total_cost": plan_cost + emergency_cost,
        "charge_kwh": sum(float(r["charge_kwh"].sum()) for r in records),
        "discharge_kwh": sum(float(r["discharge_kwh"].sum()) for r in records),
        "surplus_kwh": sum(float(r["surplus_kwh"].sum()) for r in records),
        "emergency_days": sum(bool(np.any(r["emergency_kwh"] > EMERGENCY_EPS)) for r in records),
    }


def validate(bundle: dict, summary: dict) -> dict:
    price = bundle["price"]
    records = bundle["records"]
    expected_dates = [REPORT_START + timedelta(days=i) for i in range(334)]
    if [r["date"] for r in records] != expected_dates:
        raise AssertionError("结果日期不是 2025-02-01 至 2025-12-31")
    max_balance = 0.0
    max_soc_error = 0.0
    min_soc = float("inf")
    max_soc = float("-inf")
    for i, r in enumerate(records):
        if i and abs(r["soc0"] - records[i - 1]["soc24"]) > 1e-8:
            raise AssertionError("跨日 SOC 不连续")
        q = r["purchase_kwh"]
        ch = r["charge_kwh"]
        dis = r["discharge_kwh"]
        em = r["emergency_kwh"]
        surplus = r["surplus_kwh"]
        soc = r["soc_trace"]
        balance = q + r["pv_kwh"] + dis + em - r["load_kwh"] - ch - surplus
        dynamics = np.empty(T)
        dynamics[0] = r["soc0"] + ETA * ch[0] - dis[0] / ETA
        dynamics[1:] = soc[:-1] + ETA * ch[1:] - dis[1:] / ETA
        max_balance = max(max_balance, float(np.max(np.abs(balance))))
        max_soc_error = max(max_soc_error, float(np.max(np.abs(soc - dynamics))))
        min_soc = min(min_soc, float(soc.min()), r["soc0"])
        max_soc = max(max_soc, float(soc.max()), r["soc0"])
        if np.min(np.r_[q, ch, dis, em, surplus]) < -1e-8:
            raise AssertionError("存在负的能量流")
        if np.max(ch) > P_MAX * DT + 1e-8 or np.max(dis) > P_MAX * DT + 1e-8:
            raise AssertionError("充放电功率超限")
        if soc.min() < SOC_MIN - 1e-8 or soc.max() > SOC_MAX + 1e-8:
            raise AssertionError("SOC 越界")
        if np.max(ch * dis) > 1e-8:
            raise AssertionError("同一时段同时充放电")
    if max_balance > 1e-8 or max_soc_error > 1e-8:
        raise AssertionError(f"平衡误差 {max_balance}, SOC误差 {max_soc_error}")
    recomputed = summarize(bundle)
    if abs(recomputed["total_cost"] - summary["total_cost"]) > 1e-7:
        raise AssertionError("费用汇总不一致")
    return {
        "status": "PASS",
        "intervals": len(records) * T,
        "max_balance_error_kwh": max_balance,
        "max_soc_error_kwh": max_soc_error,
        "soc_range_kwh": [min_soc, max_soc],
        "checks": [
            "date coverage",
            "continuous cross-day SOC",
            "energy balance",
            "SOC bounds",
            "power bounds",
            "no simultaneous charge and discharge",
            "cost reconciliation",
        ],
    }


def _write_jsonable_record(record: dict) -> dict:
    return {
        "date": record["date"].isoformat(),
        "soc0": record["soc0"],
        "soc24": record["soc24"],
        "purchase_kwh": record["purchase_kwh"].tolist(),
        "charge_kwh": record["charge_kwh"].tolist(),
        "discharge_kwh": record["discharge_kwh"].tolist(),
        "emergency_kwh": record["emergency_kwh"].tolist(),
        "surplus_kwh": record["surplus_kwh"].tolist(),
        "soc_trace": record["soc_trace"].tolist(),
        "load_kwh": record["load_kwh"].tolist(),
        "pv_kwh": record["pv_kwh"].tolist(),
        "reference_soc": record["reference_soc"].tolist(),
    }


def save_data(
    bundle: dict,
    summary: dict,
    validation: dict,
    parameters: Parameters,
    selection: dict,
    causality: dict,
) -> Path:
    out = ROOT / "q2_rebuild_data.json"
    payload = {
        "parameters": asdict(parameters),
        "parameter_selection": selection,
        "causality": causality,
        "summary": summary,
        "validation": validation,
        "price": bundle["price"].tolist(),
        "records": [_write_jsonable_record(r) for r in bundle["records"]],
    }
    out.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return out


def save_result2(bundle: dict) -> Path:
    """Fill the supplied result2 template for the official 2-12 month range."""
    wb = load_workbook(TEMPLATE)
    records = bundle["records"]
    record_by_date = {r["date"]: r for r in records}

    ws_plan = wb["计划购电量"]
    for t in range(T):
        ws_plan.cell(1, 2 + t, interval_label(t, t + 1))
    template_rows = {}
    for row in range(2, ws_plan.max_row + 1):
        value = ws_plan.cell(row, 1).value
        if isinstance(value, datetime):
            template_rows[value.date()] = row
        elif isinstance(value, date):
            template_rows[value] = row
    for day, record in record_by_date.items():
        row = template_rows[day]
        for t, quantity in enumerate(record["purchase_kwh"]):
            ws_plan.cell(row, 2 + t, round(float(quantity), 6))
        ws_plan.cell(row, 146, round(float(record["purchase_kwh"].sum()), 6))
        plan_cost = float(np.dot(bundle["price"], record["purchase_kwh"]))
        emergency_cost = float(np.dot(5.0 * bundle["price"], record["emergency_kwh"]))
        ws_plan.cell(row, 147, round(plan_cost + emergency_cost, 6))
    ws_plan["A2"].number_format = "yyyy/mm/dd"
    for row in range(3, 336):
        ws_plan.cell(row, 1).number_format = "yyyy/mm/dd"

    ws_storage = wb["充放电量"]
    ws_storage.delete_rows(2, ws_storage.max_row)
    bins = source_four_hour_bins()
    windows = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    row = 2
    for record in records:
        for block, (window, indices) in enumerate(zip(windows, bins)):
            ws_storage.cell(row, 1, datetime.combine(record["date"], datetime.min.time()) if block == 0 else None)
            ws_storage.cell(row, 2, window)
            ws_storage.cell(row, 3, round(float(record["charge_kwh"][indices].sum()), 6))
            ws_storage.cell(row, 4, round(float(record["discharge_kwh"][indices].sum()), 6))
            if block == 0:
                ws_storage.cell(row, 5, "0:00")
                ws_storage.cell(row, 6, round(float(record["soc0"]), 6))
            elif block == 1:
                ws_storage.cell(row, 5, "24:00")
                ws_storage.cell(row, 6, round(float(record["soc24"]), 6))
            row += 1
    for row in range(2, 2006):
        ws_storage.cell(row, 1).number_format = "yyyy/mm/dd"

    ws_emergency = wb["紧急购电量"]
    ws_emergency.delete_rows(2, ws_emergency.max_row)
    row = 2
    for record in records:
        events = emergency_segments(record["emergency_kwh"])
        if not events:
            ws_emergency.cell(row, 1, datetime.combine(record["date"], datetime.min.time()))
            ws_emergency.cell(row, 2, "无")
            ws_emergency.cell(row, 3, 0.0)
            row += 1
            continue
        for index, (span, quantity) in enumerate(events):
            ws_emergency.cell(row, 1, datetime.combine(record["date"], datetime.min.time()) if index == 0 else None)
            ws_emergency.cell(row, 2, span)
            ws_emergency.cell(row, 3, round(quantity, 6))
            row += 1
    for row in range(2, 2006):
        ws_emergency.cell(row, 1).number_format = "yyyy/mm/dd"

    for ws in wb.worksheets:
        ws.freeze_panes = "B2"
        for row_cells in ws.iter_rows(min_row=2):
            for cell in row_cells:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000000"
                elif isinstance(cell.value, datetime):
                    cell.number_format = "yyyy/mm/dd"
    output = ROOT / "result2.xlsx"
    wb.save(output)
    return output


def save_specified_markdown(bundle: dict, summary: dict, validation: dict) -> Path:
    price = bundle["price"]
    by_date = {r["date"]: r for r in bundle["records"]}
    lines = [
        "# 问题2指定日期结果",
        "",
        "计划购电量、充放电量和紧急购电量单位均为 kWh，费用单位为元。",
        "一月从 2025-01-01 0:00、6000 kWh 连续模拟，2 月 1 日继承 1 月 31 日末储电量。",
        "",
    ]
    bins = source_four_hour_bins()
    windows = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    chosen_slots = [60, 72, 84, 96, 108, 120]
    for day in SELECTED_DATES:
        r = by_date[day]
        plan_cost = float(np.dot(price, r["purchase_kwh"]))
        emergency_cost = float(np.dot(5.0 * price, r["emergency_kwh"]))
        lines += [
            f"## {day.isoformat()}",
            "",
            "### 表1 计划购电",
            "",
            "| 时间段 | 购电量 |",
            "|---|---:|",
        ]
        for slot in chosen_slots:
            lines.append(f"| {interval_label(slot, slot + 1)} | {r['purchase_kwh'][slot]:.6f} |")
        lines += [
            f"| 全天购电量 | {r['purchase_kwh'].sum():.6f} |",
            f"| 全天购电费 | {plan_cost + emergency_cost:.6f} |",
            "",
            "### 表2 充放电与储电量",
            "",
            "| 时间段 | 充电量 | 放电量 |",
            "|---|---:|---:|",
        ]
        for name, ix in zip(windows, bins):
            lines.append(f"| {name} | {r['charge_kwh'][ix].sum():.6f} | {r['discharge_kwh'][ix].sum():.6f} |")
        lines += [
            f"| 0:00储电量 | {r['soc0']:.6f} | |",
            f"| 24:00储电量 | {r['soc24']:.6f} | |",
            "",
            "### 表3 紧急购电",
            "",
            "| 紧急购电时间段 | 购电量 |",
            "|---|---:|",
        ]
        events = emergency_segments(r["emergency_kwh"])
        lines += [f"| {label} | {qty:.6f} |" for label, qty in events] or ["| 无 | 0 |"]
        lines += [f"紧急购电费用：{emergency_cost:.6f}。", ""]
    lines += [
        "## 计算汇总",
        "",
        f"2-12月计划购电费：{summary['plan_cost']:.6f}。",
        f"2-12月紧急购电费：{summary['emergency_cost']:.6f}。",
        f"2-12月合计购电费：{summary['total_cost']:.6f}。",
        f"2月1日0:00储电量：{summary['state_at_report_start']:.6f}。",
        f"12月31日24:00储电量：{summary['period_end_soc']:.6f}。",
        f"验证状态：{validation['status']}。",
    ]
    out = ROOT / "问题2指定日期结果_重建版.md"
    out.write_text("\n".join(lines), encoding="utf-8")
    return out


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=ROOT)
    parser.add_argument("--no-xlsx", action="store_true")
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    parameters = Parameters()
    warmup_parameters = WARMUP_PARAMETERS
    selection = {
        "method": "lowest previously evaluated candidate adopted after retrospective comparison",
        "selected": asdict(parameters),
        "source": "q2_borrowed_algorithm_experiment.csv",
        "warmup": asdict(WARMUP_PARAMETERS),
        "warmup_forecast": "fixed weighted same-weekday load and four-day weighted PV",
        "warning": "Decisions are causal conditional on the frozen configuration. Selection followed historical/full-year comparisons; this is not an independent held-out test or a proof of global optimality.",
    }
    bundle, summary = simulate(
        parameters,
        warmup_parameters=warmup_parameters,
        terminal_soc_target=TERMINAL_SOC_TARGET,
    )
    _, baseline_summary = simulate(
        parameters,
        warmup_parameters=warmup_parameters,
        apply_error_correction=False,
        warmup_apply_error_correction=True,
        clip_negative_demand=False,
        terminal_soc_target=TERMINAL_SOC_TARGET,
    )
    baseline_summary["method"] = "same causal forecasts without error-quantile correction"
    validation = validate(bundle, summary)
    causality = causal_smoke_test(parameters)
    xlsx_path = None if args.no_xlsx else save_result2(bundle)
    data_path = save_data(bundle, summary, validation, parameters, selection, causality)
    (ROOT / "q2_baseline_summary.json").write_text(
        json.dumps(
            {"parameters": asdict(parameters), "summary": baseline_summary},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    md_path = save_specified_markdown(bundle, summary, validation)
    (ROOT / "q2_summary.json").write_text(
        json.dumps(
            {
                "parameters": asdict(parameters),
                "parameter_selection": selection,
                "causality": causality,
                "baseline": baseline_summary,
                "summary": summary,
                "validation": validation,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(json.dumps({"xlsx": str(xlsx_path) if xlsx_path else None, "data": str(data_path), "baseline": baseline_summary, "markdown": str(md_path), "parameters": asdict(parameters), "summary": summary, "validation": validation, "causality": causality}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
