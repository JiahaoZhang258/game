"""CUMCM 2026 C题 Q2：风险敏感情景计划 + 日内因果兑现。

口径：
- 电价附件1；负荷/光伏附件2 实际值；不用附件3、不用当天实际做计划
- 日前：最近至多 4 个同星期负荷 × 近日历日光伏，近大远小加权笛卡尔积
- 日前计划可用 CVaR 对高紧急购电情景加权
- 计划购电 here-and-now；计划电必须吃进
- 日内兑现因果滚动：每个 10 分钟只用当前实际、已观测偏差和历史预报
- 1 月 1 日 0:00 SOC=6000 预热，2 月 1 日起出数；之后 SOC 不锁
- 末端库存价值 μ=0.60 元/kWh
"""
from __future__ import annotations

import json
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import coo_matrix

from solve_q1 import (
    DATA,
    DT,
    ETA,
    P_MAX,
    ROOT,
    SOC_MAX,
    SOC_MIN,
    T,
    four_hour_bins,
    load_attachment1,
)

TEMPLATE = DATA / "附件5" / "result2.xlsx"
SOC_JAN1 = 6000.0
K_MAX = 4
K_P_MAX = 4
K_MIN = 3
MU = 0.60
REG = 1e-4
EM_EPS = 1e-4
VALIDATION_TOL = 1e-4
CVaR_ALPHA = 0.75
RISK_WEIGHT_DEFAULT = 0.45
FORECAST_UPDATE_DEFAULT = "none"

# 这些验证段按时间先后排列。每一段只使用该段开始前已经积累的历史，
# 但候选算法仍从 1 月 1 日连续运行，以保留真实的储能跨日状态。
VALIDATION_FOLDS = [
    (date(2025, 3, 1), date(2025, 4, 30)),
    (date(2025, 6, 1), date(2025, 7, 31)),
    (date(2025, 8, 1), date(2025, 8, 31)),
]
HOLDOUT_WINDOW = (date(2025, 9, 1), date(2025, 12, 31))


def _excel_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return date(1899, 12, 30) + timedelta(days=int(float(v)))


def load_attachment2() -> dict:
    wb = load_workbook(DATA / "附件2.xlsx", data_only=True, read_only=True)

    def read_sheet(ws):
        rows = ws.iter_rows(values_only=True)
        next(rows)
        dates, mat = [], []
        for row in rows:
            if row[0] is None:
                continue
            dates.append(_excel_date(row[0]))
            mat.append([float(x or 0.0) for x in row[1 : 1 + T]])
        return dates, np.asarray(mat, dtype=float)

    d_load, load = read_sheet(wb["小区负载"])
    d_pv, pv = read_sheet(wb["光伏发电实际功率"])
    wb.close()
    if d_load != d_pv:
        raise RuntimeError("附件2 负荷与光伏日期不一致")
    return {"dates": d_load, "load": load, "pv": pv}


def _rank_weights(n: int, kind: str = "harmonic") -> np.ndarray:
    """下标 0 最旧、n-1 最新。harmonic: 1/n … 1；exp: 1/2^{n-1} … 1。"""
    if n <= 0:
        return np.ones(0)
    if kind == "exp":
        return np.array([0.5 ** (n - 1 - i) for i in range(n)], dtype=float)
    return np.array([1.0 / (n - i) for i in range(n)], dtype=float)


def _equal_rho(n: int) -> np.ndarray:
    return np.ones(n, dtype=float) / max(n, 1)


def _pad(loads: list, pvs: list, typical_load, typical_pv, pv_fill=None) -> tuple[np.ndarray, np.ndarray]:
    pv_fill = typical_pv if pv_fill is None else pv_fill
    if not loads:
        loads = [typical_load]
        pvs = [pv_fill]
    while len(loads) < K_MIN:
        loads.append(typical_load)
        pvs.append(pv_fill)
    return np.asarray(loads, dtype=float), np.asarray(pvs, dtype=float)


def _cartesian(loads_src, pvs_src, recency: str) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    nL, nP = len(loads_src), len(pvs_src)
    if recency in {"harmonic", "exp"}:
        wL = _rank_weights(nL, recency)
        wP = _rank_weights(nP, recency)
    else:
        wL = np.ones(nL)
        wP = np.ones(nP)
    loads, pvs, rho = [], [], []
    for i, Lrow in enumerate(loads_src):
        for j, Prow in enumerate(pvs_src):
            loads.append(Lrow)
            pvs.append(Prow)
            rho.append(float(wL[i] * wP[j]))
    rho_a = np.asarray(rho, dtype=float)
    rho_a = rho_a / rho_a.sum()
    return np.asarray(loads, dtype=float), np.asarray(pvs, dtype=float), rho_a


def persist_pv(pv_all: np.ndarray, day_index: int, typical_pv: np.ndarray) -> np.ndarray:
    if day_index <= 0:
        return typical_pv
    sl = pv_all[max(0, day_index - 3) : day_index]
    yest = pv_all[day_index - 1]
    return 0.7 * yest + 0.3 * sl.mean(axis=0)


def build_scenarios(
    mode: str,
    dates: list[date],
    load_all: np.ndarray,
    pv_all: np.ndarray,
    day_index: int,
    typical_load: np.ndarray,
    typical_pv: np.ndarray,
    k_l: int = K_MAX,
    k_p: int = K_P_MAX,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """日前情景，返回 (loads, pvs, rho)。"""
    wd = dates[day_index].weekday()
    hist = [i for i in range(day_index) if dates[i].weekday() == wd]
    if "exp" in mode:
        recency = "exp"
    elif "recency" in mode:
        recency = "harmonic"
    else:
        recency = "equal"

    if mode == "weekday_joint":
        take = hist[-k_l:]
        loads = [load_all[i] for i in take]
        pvs = [pv_all[i] for i in take]
        L, P = _pad(loads, pvs, typical_load, typical_pv)
        return L, P, _equal_rho(len(L))

    if mode == "weekday_load_persist_pv":
        take = hist[-k_l:]
        pv_hat = persist_pv(pv_all, day_index, typical_pv)
        loads = [load_all[i] for i in take]
        pvs = [pv_hat] * len(loads)
        L, P = _pad(loads, pvs, typical_load, typical_pv, pv_fill=pv_hat)
        return L, P, _equal_rho(len(L))

    if mode == "similar_weekday":
        yest = typical_pv if day_index == 0 else pv_all[day_index - 1]
        scored = []
        pool = hist[-12:] if hist else []
        for i in pool:
            prev = pv_all[i - 1] if i >= 1 else typical_pv
            scored.append((float(np.linalg.norm(prev - yest)), i))
        scored.sort()
        take = [i for _, i in scored[:k_l]]
        loads = [load_all[i] for i in take]
        pvs = [pv_all[i] for i in take]
        L, P = _pad(loads, pvs, typical_load, typical_pv)
        return L, P, _equal_rho(len(L))

    if mode.startswith("weekday_load_recent_pv"):
        take = hist[-k_l:]
        loads_src = [load_all[i] for i in take] or [typical_load]
        pv_idx = list(range(max(0, day_index - k_p), day_index))
        pvs_src = [pv_all[j] for j in pv_idx] or [typical_pv]
        L, P, rho = _cartesian(loads_src, pvs_src, recency)
        if len(L) < K_MIN:
            L2, P2 = _pad(list(L), list(P), typical_load, typical_pv)
            return L2, P2, _equal_rho(len(L2))
        return L, P, rho

    raise ValueError(mode)


def stochastic_plan(
    price,
    loads,
    pvs,
    e0: float,
    rho=None,
    mu: float = MU,
    risk_weight: float = RISK_WEIGHT_DEFAULT,
    cvar_alpha: float = CVaR_ALPHA,
) -> np.ndarray:
    """生成 here-and-now 计划购电 g（kW）。

    ``risk_weight`` 将期望紧急购电费用与 CVaR 混合。CVaR（条件风险价值）
    是最差 ``1 - cvar_alpha`` 情景的平均费用，用于压低尾部紧急购电风险。
    """
    k = loads.shape[0]
    if rho is None:
        rho = np.ones(k) / k
    else:
        rho = np.asarray(rho, dtype=float)
        rho = rho / rho.sum()
    if not 0.0 <= risk_weight <= 1.0:
        raise ValueError(f"risk_weight must be in [0, 1], got {risk_weight}")
    if not 0.0 <= cvar_alpha < 1.0:
        raise ValueError(f"cvar_alpha must be in [0, 1), got {cvar_alpha}")

    # g[T] | for ω: em, ch, dis, dump, eend each T | optional CVaR variables
    blk = 5 * T
    use_cvar = risk_weight > 0.0
    cvar_eta = T + k * blk if use_cvar else -1
    cvar_excess = cvar_eta + 1 if use_cvar else -1
    n = cvar_excess + k if use_cvar else T + k * blk
    g0 = 0

    def off(w: int) -> int:
        return T + w * blk

    cobj = np.zeros(n)
    cobj[g0 : g0 + T] = price * DT
    bounds: list[tuple[float | None, float | None]] = [(0, None)] * T
    ri, cj, data, b = [], [], [], []
    eq = 0
    ub_ri, ub_cj, ub_data, ub_b = [], [], [], []

    def add(row: int, col: int, val: float) -> None:
        ri.append(row)
        cj.append(col)
        data.append(val)

    def add_ub(row: int, col: int, val: float) -> None:
        ub_ri.append(row)
        ub_cj.append(col)
        ub_data.append(val)

    expected_em_coef = 1.0 - risk_weight
    cvar_coef = risk_weight
    if use_cvar:
        cobj[cvar_eta] = cvar_coef
        cobj[cvar_excess : cvar_excess + k] = (
            cvar_coef * rho / max(1.0 - cvar_alpha, 1e-12)
        )

    for w in range(k):
        em, ch, dis, dump, eend = (
            off(w),
            off(w) + T,
            off(w) + 2 * T,
            off(w) + 3 * T,
            off(w) + 4 * T,
        )
        cobj[em : em + T] += (5.0 * expected_em_coef * rho[w]) * price * DT
        cobj[eend + T - 1] -= mu * rho[w]
        cobj[ch : ch + T] += REG * DT
        cobj[dis : dis + T] += REG * DT
        bounds += [(0, None)] * T
        bounds += [(0, P_MAX)] * T
        bounds += [(0, P_MAX)] * T
        bounds += [(0, None)] * T
        bounds += [(SOC_MIN, SOC_MAX)] * T
        net = loads[w] - pvs[w]
        for t in range(T):
            add(eq, g0 + t, 1.0)
            add(eq, em + t, 1.0)
            add(eq, ch + t, -1.0)
            add(eq, dis + t, 1.0)
            add(eq, dump + t, -1.0)
            b.append(net[t])
            eq += 1
        for t in range(T):
            add(eq, eend + t, 1.0)
            add(eq, ch + t, -ETA * DT)
            add(eq, dis + t, DT / ETA)
            if t == 0:
                b.append(e0)
            else:
                add(eq, eend + t - 1, -1.0)
                b.append(0.0)
            eq += 1
        if use_cvar:
            # excess[w] >= emergency_cost[w] - eta
            row = w
            for t in range(T):
                add_ub(row, em + t, 5.0 * price[t] * DT)
            add_ub(row, cvar_eta, -1.0)
            add_ub(row, cvar_excess + w, -1.0)
            ub_b.append(0.0)

    if use_cvar:
        bounds += [(0, None)] * (k + 1)

    res = linprog(
        cobj,
        A_eq=coo_matrix((data, (ri, cj)), shape=(eq, n)).tocsr(),
        b_eq=np.asarray(b),
        A_ub=(
            coo_matrix(
                (ub_data, (ub_ri, ub_cj)), shape=(k, n)
            ).tocsr()
            if use_cvar
            else None
        ),
        b_ub=np.asarray(ub_b, dtype=float) if use_cvar else None,
        bounds=bounds,
        method="highs",
        options={"presolve": True},
    )
    if not res.success:
        raise RuntimeError(f"stochastic plan LP failed: {res.message}")
    return np.maximum(res.x[g0 : g0 + T], 0.0)


def execute_horizon(price, load, pv, e0: float, buy_kw: np.ndarray, mu: float = MU) -> dict:
    """剩余时段确定性 LP。load/pv 可以是实际与预报的拼接，不含未来实际。"""
    nT = int(len(price))
    n = 5 * nT + 1
    em, ch, dis, dump, eend, e0i = 0, nT, 2 * nT, 3 * nT, 4 * nT, 5 * nT
    c = np.zeros(n)
    c[em : em + nT] = 5.0 * price * DT
    c[eend + nT - 1] = -mu
    c[ch : ch + nT] += 1e-3
    c[dis : dis + nT] += 1e-3
    c[dump : dump + nT] = 1e-8
    bounds: list[tuple[float | None, float | None]] = [(0, None)] * nT
    bounds += [(0, P_MAX)] * nT
    bounds += [(0, P_MAX)] * nT
    bounds += [(0, None)] * nT
    bounds += [(SOC_MIN, SOC_MAX)] * nT
    bounds += [(e0, e0)]
    ri, cj, data, b = [], [], [], []
    eq = 0
    net = load - pv - buy_kw

    def add(row: int, col: int, val: float) -> None:
        ri.append(row)
        cj.append(col)
        data.append(val)

    for t in range(nT):
        add(eq, em + t, 1.0)
        add(eq, ch + t, -1.0)
        add(eq, dis + t, 1.0)
        add(eq, dump + t, -1.0)
        b.append(net[t])
        eq += 1
    for t in range(nT):
        add(eq, eend + t, 1.0)
        add(eq, e0i if t == 0 else eend + t - 1, -1.0)
        add(eq, ch + t, -ETA * DT)
        add(eq, dis + t, DT / ETA)
        b.append(0.0)
        eq += 1
    res = linprog(
        c,
        A_eq=coo_matrix((data, (ri, cj)), shape=(eq, n)).tocsr(),
        b_eq=np.asarray(b, dtype=float),
        bounds=bounds,
        method="highs",
        options={"presolve": True},
    )
    if not res.success:
        raise RuntimeError(f"exec LP failed: {res.message}")
    x = res.x
    return {
        "ch_kw": np.maximum(x[ch : ch + nT], 0.0),
        "dis_kw": np.maximum(x[dis : dis + nT], 0.0),
        "dump_kw": np.maximum(x[dump : dump + nT], 0.0),
        "em_kw": np.maximum(x[em : em + nT], 0.0),
        "soc_end": x[eend : eend + nT],
    }


def execute_day(price, load, pv, e0: float, buy_kw: np.ndarray, mu: float = MU) -> dict:
    """对照：一次看见全天实际（wait-and-see），不作正式成绩。"""
    return execute_horizon(price, load, pv, e0, buy_kw, mu)


def _greedy_slot(net: float, e: float) -> tuple[float, float, float, float, float]:
    """单时段因果可行动作：余电先充后弃，缺口先放后紧急。"""
    if net <= 0:
        room = (SOC_MAX - e) / (ETA * DT)
        ch = min(-net, P_MAX, max(room, 0.0))
        dump = -net - ch
        e2 = e + ETA * ch * DT
        return ch, 0.0, 0.0, dump, e2
    head = ETA * (e - SOC_MIN) / DT
    dis = min(net, P_MAX, max(head, 0.0))
    em = net - dis
    e2 = e - dis * DT / ETA
    return 0.0, dis, em, 0.0, e2


def _causal_scale(
    actual: np.ndarray,
    baseline: np.ndarray,
    end: int,
    lower: float,
    upper: float,
    min_signal: float = 1.0,
    warmup: int = 12,
    shrink: float = 1.0,
) -> float:
    """用截至 ``end`` 的观测估计比例，并逐步收缩到 1。

    采用无截距最小二乘比例
    ``sum(actual * baseline) / sum(baseline ** 2)``，只读取当前及过去
    时段。``warmup`` 用来避免刚观察到一个点时对全天预测过度反应。
    """
    actual = np.asarray(actual[:end], dtype=float)
    baseline = np.asarray(baseline[:end], dtype=float)
    valid = (
        np.isfinite(actual)
        & np.isfinite(baseline)
        & (baseline > min_signal)
    )
    if not valid.any():
        return 1.0
    base = baseline[valid]
    obs = np.maximum(actual[valid], 0.0)
    denom = float(np.dot(base, base))
    if denom <= 1e-12:
        return 1.0
    raw = float(np.dot(obs, base) / denom)
    raw = float(np.clip(raw, lower, upper))
    confidence = min(1.0, float(valid.sum()) / max(warmup, 1))
    return 1.0 + shrink * confidence * (raw - 1.0)


def _causal_forecast(
    load_act: np.ndarray,
    pv_act: np.ndarray,
    load_f: np.ndarray,
    pv_f: np.ndarray,
    t: int,
    forecast_update: str,
) -> tuple[np.ndarray, np.ndarray]:
    """返回时刻 ``t`` 可用的剩余预测，不读取 ``t`` 之后的实际值。"""
    if forecast_update not in {
        "none",
        "adaptive_ratio",
        "pv_ratio",
        "shrunk_ratio",
    }:
        raise ValueError(f"unknown forecast_update: {forecast_update}")

    Lf = np.asarray(load_f[t:], dtype=float).copy()
    Pf = np.asarray(pv_f[t:], dtype=float).copy()
    if forecast_update in {"adaptive_ratio", "shrunk_ratio"}:
        shrink = 0.5 if forecast_update == "shrunk_ratio" else 1.0
        load_lower, load_upper = (0.90, 1.12) if shrink < 1.0 else (0.85, 1.20)
        load_scale = _causal_scale(
            load_act,
            load_f,
            t + 1,
            lower=load_lower,
            upper=load_upper,
            min_signal=1.0,
            warmup=36 if shrink < 1.0 else 12,
            shrink=shrink,
        )
        pv_floor = max(1.0, 0.05 * float(np.max(pv_f)))
        pv_scale = _causal_scale(
            pv_act,
            pv_f,
            t + 1,
            lower=0.70 if shrink < 1.0 else 0.45,
            upper=1.30 if shrink < 1.0 else 1.55,
            min_signal=pv_floor,
            warmup=36 if shrink < 1.0 else 12,
            shrink=shrink,
        )
        if len(Lf) > 1:
            Lf[1:] = np.maximum(Lf[1:] * load_scale, 0.0)
            Pf[1:] = np.maximum(Pf[1:] * pv_scale, 0.0)
    elif forecast_update == "pv_ratio":
        pv_floor = max(1.0, 0.05 * float(np.max(pv_f)))
        pv_scale = _causal_scale(
            pv_act,
            pv_f,
            t + 1,
            lower=0.70,
            upper=1.30,
            min_signal=pv_floor,
            warmup=36,
            shrink=0.5,
        )
        if len(Pf) > 1:
            Pf[1:] = np.maximum(Pf[1:] * pv_scale, 0.0)

    # 当前格在求解时已观测，后续格仍只能使用更新后的预测。
    Lf[0] = max(float(load_act[t]), 0.0)
    Pf[0] = max(float(pv_act[t]), 0.0)
    return Lf, Pf


def execute_day_causal(
    price,
    load_act,
    pv_act,
    e0: float,
    buy_kw: np.ndarray,
    load_f: np.ndarray,
    pv_f: np.ndarray,
    mu: float = MU,
    forecast_update: str = FORECAST_UPDATE_DEFAULT,
) -> dict:
    """因果滚动：只执行当前格，未来实际只通过已观测偏差修正预测。"""
    e = float(e0)
    ch = np.zeros(T)
    dis = np.zeros(T)
    em = np.zeros(T)
    dump = np.zeros(T)
    soc_end = np.zeros(T)
    for t in range(T):
        Lf, Pf = _causal_forecast(
            load_act, pv_act, load_f, pv_f, t, forecast_update
        )
        try:
            sol = execute_horizon(price[t:], Lf, Pf, e, buy_kw[t:], mu)
            ch[t] = float(sol["ch_kw"][0])
            dis[t] = float(sol["dis_kw"][0])
            em[t] = float(sol["em_kw"][0])
            dump[t] = float(sol["dump_kw"][0])
            e = float(np.clip(sol["soc_end"][0], SOC_MIN, SOC_MAX))
        except RuntimeError:
            net = float(load_act[t] - pv_act[t] - buy_kw[t])
            ch[t], dis[t], em[t], dump[t], e = _greedy_slot(net, e)
            e = float(np.clip(e, SOC_MIN, SOC_MAX))
        soc_end[t] = e
    return {
        "ch_kw": ch,
        "dis_kw": dis,
        "dump_kw": dump,
        "em_kw": em,
        "soc_end": soc_end,
    }


def _period_cost_summary(records: list[dict], start: date, end: date) -> dict:
    selected = [r for r in records if start <= r["date"] <= end]
    if not selected:
        raise ValueError(f"no records in validation window {start}..{end}")
    return {
        "start": str(start),
        "end": str(end),
        "n_days": len(selected),
        "total_cost": float(sum(r["total_cost"] for r in selected)),
        "plan_cost": float(sum(r["plan_cost"] for r in selected)),
        "em_cost": float(sum(r["em_cost"] for r in selected)),
        "em_kwh": float(sum(r["em_kwh"].sum() for r in selected)),
    }


def rolling_validation(records: list[dict]) -> dict:
    """按时间顺序的固定验证段汇总，并保留最终留出段。"""
    folds = [
        _period_cost_summary(records, start, end)
        for start, end in VALIDATION_FOLDS
    ]
    holdout = _period_cost_summary(records, *HOLDOUT_WINDOW)
    total_days = sum(f["n_days"] for f in folds)
    total_cost = sum(f["total_cost"] for f in folds)
    return {
        "folds": folds,
        "total_days": total_days,
        "total_cost": float(total_cost),
        "mean_daily_cost": float(total_cost / max(total_days, 1)),
        "holdout": holdout,
        "selection_metric": "rolling_validation_total_cost",
    }


def merge_emergency(minutes: np.ndarray, em_kwh: np.ndarray) -> list[tuple[str, float]]:
    segs = []
    i = 0
    while i < T:
        if em_kwh[i] <= EM_EPS:
            i += 1
            continue
        j = i
        acc = 0.0
        while j < T and em_kwh[j] > EM_EPS:
            acc += em_kwh[j]
            j += 1
        start = int(minutes[i])
        end = int(minutes[j - 1]) + 10
        segs.append((_span_label(start, end), float(acc)))
        i = j
    return segs


def _span_label(start_min: int, end_min: int) -> str:
    def fmt(m: int) -> str:
        if m >= 24 * 60:
            mm = m - 24 * 60
            h, mi = divmod(mm, 60)
            return f"{h}:{mi:02d}+1" if mm else "0:00+1"
        h, mi = divmod(m, 60)
        return f"{h}:{mi:02d}"

    return f"{fmt(start_min)}-{fmt(end_min)}"


def run(
    mode: str = "weekday_joint",
    quiet: bool = False,
    k_l: int = K_MAX,
    k_p: int = K_P_MAX,
    mu: float = MU,
    risk_weight: float = RISK_WEIGHT_DEFAULT,
    cvar_alpha: float = CVaR_ALPHA,
    forecast_update: str = FORECAST_UPDATE_DEFAULT,
) -> dict:
    att1 = load_attachment1(DATA / "附件1.xlsx")
    att2 = load_attachment2()
    price, minutes = att1["price"], att1["minutes"]
    typical_load, typical_pv = att1["load"], att1["pv"]
    dates = att2["dates"]
    load_all, pv_all = att2["load"], att2["pv"]

    report_from = date(2025, 2, 1)
    soc = SOC_JAN1
    records = []
    soc_feb1 = None

    for i, d in enumerate(dates):
        loads, pvs, rho = build_scenarios(
            mode, dates, load_all, pv_all, i, typical_load, typical_pv, k_l=k_l, k_p=k_p
        )
        buy_kw = stochastic_plan(
            price,
            loads,
            pvs,
            soc,
            rho=rho,
            mu=mu,
            risk_weight=risk_weight,
            cvar_alpha=cvar_alpha,
        )
        load_f = rho @ loads
        pv_f = rho @ pvs
        exe = execute_day_causal(
            price,
            load_all[i],
            pv_all[i],
            soc,
            buy_kw,
            load_f,
            pv_f,
            mu=mu,
            forecast_update=forecast_update,
        )
        exe_ws = execute_day(price, load_all[i], pv_all[i], soc, buy_kw, mu=mu)
        buy_kwh = buy_kw * DT
        em_kwh = np.maximum(exe["em_kw"] * DT, 0.0)
        em_ws = np.maximum(exe_ws["em_kw"] * DT, 0.0)
        rec = {
            "date": d,
            "n_scen": int(loads.shape[0]),
            "soc0": soc,
            "soc24": float(exe["soc_end"][-1]),
            "soc_trace": np.asarray(exe["soc_end"], dtype=float),
            "soc_min": float(min(soc, exe["soc_end"].min())),
            "soc_max": float(max(soc, exe["soc_end"].max())),
            "load_kwh": load_all[i] * DT,
            "pv_kwh": pv_all[i] * DT,
            "buy_kwh": buy_kwh,
            "ch_kwh": exe["ch_kw"] * DT,
            "dis_kwh": exe["dis_kw"] * DT,
            "em_kwh": em_kwh,
            "dump_kwh": exe["dump_kw"] * DT,
            "plan_cost": float(np.dot(price, buy_kwh)),
            "em_cost": float(np.dot(price, 5.0 * em_kwh)),
            "ws_em_kwh": float(em_ws.sum()),
            "ws_cost": float(np.dot(price, buy_kwh) + np.dot(price, 5.0 * em_ws)),
        }
        rec["total_cost"] = rec["plan_cost"] + rec["em_cost"]
        rec["em_segs"] = merge_emergency(minutes, em_kwh)
        if d == report_from:
            soc_feb1 = soc
        if d >= report_from:
            records.append(rec)
        soc = rec["soc24"]
        if (not quiet) and (
            d.day == 1
            or d == dates[0]
            or d == dates[-1]
            or d
            in {
                date(2025, 2, 1),
                date(2025, 3, 20),
                date(2025, 6, 21),
                date(2025, 9, 23),
                date(2025, 12, 21),
            }
        ):
            print(
                f"{d} wd={d.weekday()} K={loads.shape[0]} "
                f"soc {rec['soc0']:.1f}->{rec['soc24']:.1f} "
                f"buy {buy_kwh.sum():.0f} em {em_kwh.sum():.1f} "
                f"cost {rec['total_cost']:.1f}",
                flush=True,
            )

    summary = {
        "mode": mode,
        "mu": mu,
        "risk_weight": risk_weight,
        "cvar_alpha": cvar_alpha,
        "forecast_update": forecast_update,
        "k_l": k_l,
        "k_p": k_p,
        "soc_jan1": SOC_JAN1,
        "soc_feb1": soc_feb1,
        "n_days": len(records),
        "total_plan_kwh": float(sum(r["buy_kwh"].sum() for r in records)),
        "total_em_kwh": float(sum(r["em_kwh"].sum() for r in records)),
        "total_ws_em_kwh": float(sum(r["ws_em_kwh"] for r in records)),
        "total_plan_cost": float(sum(r["plan_cost"] for r in records)),
        "total_em_cost": float(sum(r["em_cost"] for r in records)),
        "total_cost": float(sum(r["total_cost"] for r in records)),
        "total_ws_cost": float(sum(r["ws_cost"] for r in records)),
        "days_with_em": int(sum(1 for r in records if r["em_kwh"].sum() > EM_EPS)),
        "execution": "causal_mpc",
        "soc_end_year": float(records[-1]["soc24"]),
        "midnight_soc_min": float(min(min(r["soc0"], r["soc24"]) for r in records)),
        "midnight_soc_max": float(max(max(r["soc0"], r["soc24"]) for r in records)),
        "intraday_soc_min": float(min(r["soc_min"] for r in records)),
        "intraday_soc_max": float(max(r["soc_max"] for r in records)),
        "max_simultaneous_kw": float(
            max(np.minimum(r["ch_kwh"] / DT, r["dis_kwh"] / DT).max() for r in records)
        ),
        "mean_dump_kwh": float(np.mean([r["dump_kwh"].sum() for r in records])),
    }
    summary["rolling_validation"] = rolling_validation(records)
    bundle = {"minutes": minutes, "price": price, "records": records, "summary": summary}
    bundle["validation"] = validate_bundle(bundle)
    return bundle


def validate_bundle(bundle: dict) -> dict:
    """校验正式因果结果的日期、能量平衡、SOC 和费用分解。"""
    records = bundle["records"]
    price = np.asarray(bundle["price"], dtype=float)
    if len(price) != T:
        raise AssertionError(f"price length is {len(price)}, expected {T}")

    expected_dates = [date(2025, 2, 1) + timedelta(days=i) for i in range(334)]
    actual_dates = [r["date"] for r in records]
    if actual_dates != expected_dates:
        raise AssertionError(
            f"result dates do not cover 2025-02-01..2025-12-31: "
            f"{actual_dates[:1]}..{actual_dates[-1:]}"
        )

    max_power_balance_error = 0.0
    max_soc_error = 0.0
    max_soc = -np.inf
    min_soc = np.inf
    max_simultaneous_kw = 0.0
    for i, rec in enumerate(records):
        arrays = {
            "buy_kwh": rec["buy_kwh"],
            "ch_kwh": rec["ch_kwh"],
            "dis_kwh": rec["dis_kwh"],
            "em_kwh": rec["em_kwh"],
            "dump_kwh": rec["dump_kwh"],
            "soc_trace": rec["soc_trace"],
        }
        for name, arr in arrays.items():
            arr = np.asarray(arr, dtype=float)
            if arr.shape != (T,):
                raise AssertionError(f"{rec['date']} {name} shape is {arr.shape}, expected {(T,)}")
            if not np.isfinite(arr).all():
                raise AssertionError(f"{rec['date']} {name} contains non-finite values")
            if name != "soc_trace" and arr.min() < -VALIDATION_TOL:
                raise AssertionError(f"{rec['date']} {name} contains negative values")

        ch = np.asarray(rec["ch_kwh"], dtype=float)
        dis = np.asarray(rec["dis_kwh"], dtype=float)
        if max(ch.max(), dis.max()) > P_MAX * DT + VALIDATION_TOL:
            raise AssertionError(f"{rec['date']} charge/discharge exceeds power limit")
        trace = np.asarray(rec["soc_trace"], dtype=float)
        max_soc = max(max_soc, float(trace.max()), float(rec["soc0"]))
        min_soc = min(min_soc, float(trace.min()), float(rec["soc0"]))
        if trace.min() < SOC_MIN - VALIDATION_TOL or trace.max() > SOC_MAX + VALIDATION_TOL:
            raise AssertionError(f"{rec['date']} SOC is outside [{SOC_MIN}, {SOC_MAX}]")
        if i > 0 and abs(float(rec["soc0"]) - float(records[i - 1]["soc24"])) > VALIDATION_TOL:
            raise AssertionError(f"{rec['date']} SOC0 is not previous day's SOC24")

        expected_trace = np.empty(T, dtype=float)
        expected_trace[0] = rec["soc0"] + ETA * ch[0] - dis[0] / ETA
        expected_trace[1:] = trace[:-1] + ETA * ch[1:] - dis[1:] / ETA
        soc_error = float(np.max(np.abs(trace - expected_trace)))
        max_soc_error = max(max_soc_error, soc_error)
        if soc_error > VALIDATION_TOL:
            raise AssertionError(f"{rec['date']} SOC dynamics error is {soc_error}")
        if abs(float(rec["soc24"]) - float(trace[-1])) > VALIDATION_TOL:
            raise AssertionError(f"{rec['date']} SOC24 does not match SOC trace")

        # All terms are already energy quantities in kWh.
        actual_load = np.asarray(rec["load_kwh"], dtype=float)
        actual_pv = np.asarray(rec["pv_kwh"], dtype=float)
        lhs = rec["buy_kwh"] + actual_pv + dis + rec["em_kwh"]
        rhs = actual_load + ch + rec["dump_kwh"]
        power_error = float(np.max(np.abs(lhs - rhs)))
        max_power_balance_error = max(max_power_balance_error, power_error)
        if power_error > VALIDATION_TOL:
            raise AssertionError(f"{rec['date']} power balance error is {power_error}")

        simultaneous = float(np.minimum(ch, dis).max() / DT)
        max_simultaneous_kw = max(max_simultaneous_kw, simultaneous)
        if simultaneous > VALIDATION_TOL:
            raise AssertionError(f"{rec['date']} charges and discharges simultaneously")

        plan_cost = float(np.dot(price, rec["buy_kwh"]))
        em_cost = float(np.dot(price, 5.0 * rec["em_kwh"]))
        if abs(plan_cost - rec["plan_cost"]) > VALIDATION_TOL:
            raise AssertionError(f"{rec['date']} plan cost mismatch")
        if abs(em_cost - rec["em_cost"]) > VALIDATION_TOL:
            raise AssertionError(f"{rec['date']} emergency cost mismatch")
        if abs(rec["total_cost"] - plan_cost - em_cost) > VALIDATION_TOL:
            raise AssertionError(f"{rec['date']} total cost mismatch")

    summary = bundle["summary"]
    totals = {
        "total_plan_kwh": sum(float(r["buy_kwh"].sum()) for r in records),
        "total_em_kwh": sum(float(r["em_kwh"].sum()) for r in records),
        "total_plan_cost": sum(float(r["plan_cost"]) for r in records),
        "total_em_cost": sum(float(r["em_cost"]) for r in records),
        "total_cost": sum(float(r["total_cost"]) for r in records),
    }
    for name, value in totals.items():
        if abs(float(summary[name]) - value) > VALIDATION_TOL:
            raise AssertionError(f"summary {name} mismatch")

    return {
        "status": "passed",
        "evaluation": "causal_mpc",
        "n_days": len(records),
        "date_range": [str(expected_dates[0]), str(expected_dates[-1])],
        "max_power_balance_error_kwh": max_power_balance_error,
        "max_soc_dynamics_error_kwh": max_soc_error,
        "soc_range_kwh": [min_soc, max_soc],
        "max_simultaneous_kw": max_simultaneous_kw,
        "checked_costs": True,
    }


def write_result2(bundle: dict) -> Path:
    minutes = bundle["minutes"]
    records = bundle["records"]
    wb = load_workbook(TEMPLATE)
    ws = wb["计划购电量"]
    date_row = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, datetime):
            date_row[v.date()] = r
        elif isinstance(v, date):
            date_row[v] = r

    for rec in records:
        r = date_row[rec["date"]]
        for t, val in enumerate(rec["buy_kwh"]):
            ws.cell(r, 2 + t, round(float(val), 4))
        ws.cell(r, 146, round(float(rec["buy_kwh"].sum()), 4))
        ws.cell(r, 147, round(float(rec["total_cost"]), 4))

    windows = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    bins = four_hour_bins(minutes)
    ws2 = wb["充放电量"]
    ws2.delete_rows(2, ws2.max_row)
    row = 2
    for rec in records:
        for k, (name, ix) in enumerate(zip(windows, bins)):
            ws2.cell(row, 1, datetime(rec["date"].year, rec["date"].month, rec["date"].day) if k == 0 else None)
            ws2.cell(row, 2, name)
            ws2.cell(row, 3, round(float(rec["ch_kwh"][ix].sum()), 4))
            ws2.cell(row, 4, round(float(rec["dis_kwh"][ix].sum()), 4))
            if k == 0:
                ws2.cell(row, 5, datetime(1900, 1, 1, 0, 0).time())
                ws2.cell(row, 6, round(float(rec["soc0"]), 4))
            elif k == 1:
                ws2.cell(row, 5, "24:00")
                ws2.cell(row, 6, round(float(rec["soc24"]), 4))
            row += 1

    ws3 = wb["紧急购电量"]
    ws3.delete_rows(2, ws3.max_row)
    row = 2
    for rec in records:
        segs = rec["em_segs"]
        if not segs:
            ws3.cell(row, 1, datetime(rec["date"].year, rec["date"].month, rec["date"].day))
            ws3.cell(row, 2, "无")
            ws3.cell(row, 3, 0)
            row += 1
            continue
        for j, (span, qty) in enumerate(segs):
            if j == 0:
                ws3.cell(row, 1, datetime(rec["date"].year, rec["date"].month, rec["date"].day))
            ws3.cell(row, 2, span)
            ws3.cell(row, 3, round(qty, 4))
            row += 1

    out = ROOT / "result2.xlsx"
    wb.save(out)
    return out


def specified_tables(bundle: dict) -> dict:
    want = [date(2025, 3, 20), date(2025, 6, 21), date(2025, 9, 23), date(2025, 12, 21)]
    minutes = bundle["minutes"]
    slots = {
        "10:00-10:10": 10 * 60,
        "12:00-12:10": 12 * 60,
        "14:00-14:10": 14 * 60,
        "16:00-16:10": 16 * 60,
        "18:00-18:10": 18 * 60,
        "20:00-20:10": 20 * 60,
    }
    windows = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    bins = four_hour_bins(minutes)
    by_date = {r["date"]: r for r in bundle["records"]}
    out = {}
    for d in want:
        rec = by_date[d]
        t1 = {}
        for name, m in slots.items():
            idx = int(np.where(minutes == m)[0][0])
            t1[name] = float(rec["buy_kwh"][idx])
        t1["全天购电量"] = float(rec["buy_kwh"].sum())
        t1["全天购电费"] = float(rec["total_cost"])
        t2 = {}
        for name, ix in zip(windows, bins):
            t2[name] = {
                "charge": float(rec["ch_kwh"][ix].sum()),
                "discharge": float(rec["dis_kwh"][ix].sum()),
            }
        t2["soc0"] = rec["soc0"]
        t2["soc24"] = rec["soc24"]
        out[str(d)] = {
            "table1": t1,
            "table2": t2,
            "emergency": rec["em_segs"],
            "em_kwh": float(rec["em_kwh"].sum()),
            "dump_kwh": float(rec["dump_kwh"].sum()),
        }
    return out


OFFICIAL = {
    "mode": "weekday_load_recent_pv_exp",
    "k_l": 4,
    "k_p": 4,
    "mu": 0.60,
    "risk_weight": 0.45,
    "cvar_alpha": CVaR_ALPHA,
    "forecast_update": "none",
    "tag": "exp_4x4_r45_none",
}

JOBS = [
    {
        "mode": "weekday_load_recent_pv_exp",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.00,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "none",
        "tag": "exp_4x4_mu60",
    },
    {
        "mode": "weekday_load_recent_pv_exp",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.15,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "none",
        "tag": "exp_4x4_r15_none",
    },
    {
        "mode": "weekday_load_recent_pv_exp",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.30,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "none",
        "tag": "exp_4x4_r30_none",
    },
    {
        "mode": "weekday_load_recent_pv_exp",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.45,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "none",
        "tag": "exp_4x4_r45_none",
    },
    {
        "mode": "weekday_load_recent_pv_exp",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.00,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "adaptive_ratio",
        "tag": "exp_4x4_r0_adapt",
    },
    {
        "mode": "weekday_load_recent_pv_exp",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.15,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "adaptive_ratio",
        "tag": "exp_4x4_r15_adapt",
    },
    {
        "mode": "weekday_load_recent_pv_exp",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.30,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "adaptive_ratio",
        "tag": "exp_4x4_r30_adapt",
    },
    {
        "mode": "weekday_load_recent_pv_exp",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.45,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "adaptive_ratio",
        "tag": "exp_4x4_r45_adapt",
    },
    {
        "mode": "weekday_load_recent_pv_recency",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.00,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "adaptive_ratio",
        "tag": "recency_4x4_r0_adapt",
    },
    {
        "mode": "weekday_load_recent_pv_recency",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.15,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "adaptive_ratio",
        "tag": "recency_4x4_r15_adapt",
    },
    {
        "mode": "weekday_load_recent_pv_recency",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.30,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "adaptive_ratio",
        "tag": "recency_4x4_r30_adapt",
    },
    {
        "mode": "weekday_load_recent_pv_recency",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.45,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "adaptive_ratio",
        "tag": "recency_4x4_r45_adapt",
    },
    {
        "mode": "weekday_load_recent_pv_exp",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.30,
        "cvar_alpha": 0.85,
        "forecast_update": "adaptive_ratio",
        "tag": "exp_4x4_r30_a85_adapt",
    },
    {
        "mode": "weekday_load_recent_pv_recency",
        "k_l": 4,
        "k_p": 6,
        "mu": 0.60,
        "risk_weight": 0.00,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "none",
        "tag": "recency_4x6_r0_none",
    },
    {
        "mode": "weekday_load_recent_pv_recency",
        "k_l": 4,
        "k_p": 5,
        "mu": 0.60,
        "risk_weight": 0.00,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "none",
        "tag": "recency_4x5_r0_none",
    },
    {
        "mode": "weekday_load_recent_pv_recency",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.00,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "none",
        "tag": "recency_4x4_r0_none",
    },
    {
        "mode": "weekday_load_recent_pv_exp",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.45,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "pv_ratio",
        "tag": "exp_4x4_r45_pv",
    },
    {
        "mode": "weekday_load_recent_pv_exp",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.45,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "shrunk_ratio",
        "tag": "exp_4x4_r45_shrunk",
    },
    {
        "mode": "weekday_load_recent_pv_exp",
        "k_l": 4,
        "k_p": 4,
        "mu": 0.60,
        "risk_weight": 0.30,
        "cvar_alpha": CVaR_ALPHA,
        "forecast_update": "shrunk_ratio",
        "tag": "exp_4x4_r30_shrunk",
    },
]


def _job_by_tag(tag: str) -> dict:
    for job in JOBS:
        if job["tag"] == tag:
            return job
    if OFFICIAL["tag"] == tag:
        return OFFICIAL
    raise ValueError(f"unknown job tag: {tag}")


def _run_job(job: dict) -> tuple[str, dict, dict]:
    tag = job["tag"]
    print(
        f"START {tag} mode={job['mode']} K_L={job['k_l']} K_P={job['k_p']} "
        f"mu={job['mu']} risk={job.get('risk_weight', 0.0)} "
        f"update={job.get('forecast_update', 'none')}",
        flush=True,
    )
    bundle = run(
        job["mode"],
        quiet=True,
        k_l=job["k_l"],
        k_p=job["k_p"],
        mu=job["mu"],
        risk_weight=job.get("risk_weight", RISK_WEIGHT_DEFAULT),
        cvar_alpha=job.get("cvar_alpha", CVaR_ALPHA),
        forecast_update=job.get("forecast_update", FORECAST_UPDATE_DEFAULT),
    )
    s = bundle["summary"]
    rv = s["rolling_validation"]
    row = {
        "tag": tag,
        "mode": job["mode"],
        "k_l": job["k_l"],
        "k_p": job["k_p"],
        "mu": job["mu"],
        "risk_weight": job.get("risk_weight", RISK_WEIGHT_DEFAULT),
        "cvar_alpha": job.get("cvar_alpha", CVaR_ALPHA),
        "forecast_update": job.get("forecast_update", FORECAST_UPDATE_DEFAULT),
        "evaluation": s["execution"],
        "selection_metric": rv["selection_metric"],
        "selection_cost": rv["total_cost"],
        "selection_mean_daily_cost": rv["mean_daily_cost"],
        "selection_days": rv["total_days"],
        "holdout_cost": rv["holdout"]["total_cost"],
        "holdout_days": rv["holdout"]["n_days"],
        "total_cost": s["total_cost"],
        "plan_cost": s["total_plan_cost"],
        "em_cost": s["total_em_cost"],
        "ws_cost": s["total_ws_cost"],
        "plan_kwh": s["total_plan_kwh"],
        "em_kwh": s["total_em_kwh"],
        "ws_em_kwh": s["total_ws_em_kwh"],
        "days_with_em": s["days_with_em"],
        "mean_dump": s["mean_dump_kwh"],
        "soc_feb1": s["soc_feb1"],
        "soc_end_year": s["soc_end_year"],
        "validation": bundle["validation"]["status"],
    }
    print(
        f"DONE {tag}: select={rv['total_cost']/1e4:.2f}万  "
        f"full={s['total_cost']/1e4:.2f}万  ws={s.get('total_ws_cost',0)/1e4:.2f}万  "
        f"em={s['total_em_kwh']:.0f}  days={s['days_with_em']} "
        f"dump={s['mean_dump_kwh']:.0f}",
        flush=True,
    )
    return tag, row, bundle


def _json_tables(tables: dict) -> dict:
    return {
        k: {
            "table1": v["table1"],
            "table2": v["table2"],
            "emergency": v["emergency"],
            "em_kwh": v["em_kwh"],
            "dump_kwh": v["dump_kwh"],
        }
        for k, v in tables.items()
    }


def _save_outputs(bundle: dict, cmp_rows: list, selected_tag: str) -> None:
    out = write_result2(bundle)
    tables = specified_tables(bundle)
    selected_row = next((r for r in cmp_rows if r.get("tag") == selected_tag), None)
    slim = {
        "compare": cmp_rows,
        "official_tag": selected_tag,
        "selected_tag": selected_tag,
        "selected_row": selected_row,
        "selection_metric": "rolling_validation_total_cost",
        "evaluation": "causal_mpc",
        "summary": bundle["summary"],
        "validation": bundle["validation"],
        "specified": _json_tables(tables),
        "result_file": str(out),
    }
    (ROOT / "q2_summary.json").write_text(json.dumps(slim, ensure_ascii=False, indent=2), encoding="utf-8")
    print("==== specified dates ====")
    print(json.dumps(tables, ensure_ascii=False, indent=2))
    print("Wrote", out, "selected=", selected_tag)


def main() -> None:
    import sys

    grid = "--grid" in sys.argv
    tag_arg = next((arg for arg in sys.argv[1:] if arg.startswith("--tag=")), None)
    if grid:
        cmp_rows = []
        bundles = {}
        with ProcessPoolExecutor(max_workers=min(len(JOBS), 8)) as pool:
            futs = {pool.submit(_run_job, job): job["tag"] for job in JOBS}
            for fut in as_completed(futs):
                tag, row, bundle = fut.result()
                bundles[tag] = bundle
                cmp_rows.append(row)
        cmp_rows.sort(key=lambda r: (r["selection_cost"], r["total_cost"]))
        selected_tag = cmp_rows[0]["tag"]
        print("\n==== COMPARE (rolling validation folds) ====")
        print(json.dumps(cmp_rows, ensure_ascii=False, indent=2))
        print("GRID_BEST", selected_tag, "metric=rolling_validation_total_cost")
        bundle = bundles[selected_tag]
        _save_outputs(bundle, cmp_rows, selected_tag)
        return

    job = _job_by_tag(tag_arg.removeprefix("--tag=")) if tag_arg else OFFICIAL
    tag, row, bundle = _run_job(job)
    cmp_rows = [row]
    print(
        "SELECTED",
        tag,
        f"select={row['selection_cost']/1e4:.2f}万",
        f"full={row['total_cost']/1e4:.2f}万",
        "metric=rolling_validation_total_cost",
    )
    _save_outputs(bundle, cmp_rows, tag)


if __name__ == "__main__":
    main()
