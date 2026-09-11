"""Question 3/4 causal rolling strategies.

Question 3 uses attachment 3 PV forecasts released at 00:00, 06:00,
12:00 and 18:00.  At each release only the not-yet-executed purchase
schedule is replanned.  Question 4 uses the same policy with attachment 4
prices.  January is simulated to carry the physical SOC into February.
"""
from __future__ import annotations

import json
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import coo_matrix

from solve_q1 import DATA, DT, ETA, P_MAX, ROOT, SOC_MAX, SOC_MIN, T, four_hour_bins, load_attachment1
from solve_q2 import (
    EMERGENCY_EPS,
    INITIAL_SOC,
    Parameters as Q2Parameters,
    WARMUP_PARAMETERS,
    design_demand,
    execute_day,
    interval_label,
    forecast_day,
    load_attachment2,
    plan_day,
)

REPORT_START = date(2025, 2, 1)
REPORT_DAYS = 334
UPDATES = (36, 72, 108)
BLEND = 0.5
ALPHA = 0.7
MU = 0.30
RESERVE_RATIO = 0.625
TERMINAL_SOC_TARGET = 2400.0
REGULARIZER = 1e-7


@dataclass(frozen=True)
class Q34Parameters:
    load_error_window: int = 7
    adjustment_times: tuple[int, ...] = UPDATES
    blend: float = BLEND
    alpha: float = ALPHA
    terminal_value: float = MU
    reserve_ratio: float = RESERVE_RATIO
    terminal_soc_target: float = TERMINAL_SOC_TARGET
    price_window: int = 4
    price_method: str = "weighted"


def _as_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date(1899, 12, 30) + timedelta(days=int(float(value)))


def load_attachment3() -> dict[date, dict[int, np.ndarray]]:
    wb = load_workbook(DATA / "附件3.xlsx", read_only=True, data_only=True)
    ws = wb.active
    out: dict[date, dict[int, np.ndarray]] = {}
    current: date | None = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] not in (None, ""):
            raw = str(row[0])
            y, m, d = [int(x) for x in raw.split("-")]
            current = date(y, m, d)
            out[current] = {}
        if current is None:
            raise ValueError("附件3首行没有日期")
        hour = int(str(row[1]).split(":")[0])
        values = np.asarray([float(x or 0.0) for x in row[2:26]], dtype=float)
        if values.shape != (24,):
            raise ValueError(f"附件3 {current} {hour}:00 不是24小时预报")
        out[current][hour] = np.maximum(values, 0.0)
    wb.close()
    if len(out) != 365 or any(set(v) != {0, 6, 12, 18} for v in out.values()):
        raise ValueError("附件3日期或发布时刻不完整")
    return out


def pv_forecast_for_day(
    forecasts: dict[int, np.ndarray], issue_hour: int
) -> np.ndarray:
    """Interpolate hourly forecast bins onto 10-minute slots of the same day."""
    values = forecasts[issue_hour]
    # The source labels these columns "forecast 1 hour" through
    # "forecast 24 hours", so the first value belongs one hour after the
    # release time. Interpolate at the same interval endpoints as the
    # observations; hold the first forecast before the first forecast hour.
    grid = np.arange(1, 25, dtype=float)
    endpoint_hours = (np.arange(T, dtype=float) + 1.0) / 6.0
    rel_slots = endpoint_hours - issue_hour
    result = np.interp(rel_slots, grid, values, left=values[0], right=values[-1])
    result[endpoint_hours <= issue_hour] = 0.0
    return np.maximum(result, 0.0)


def load_history_forecast(day_index: int, dates, load, typical_load) -> np.ndarray:
    """Causal mean of the last four completed same-weekday load rows."""
    if day_index == 0:
        source = typical_load
    else:
        history = [load[i] for i in range(day_index) if dates[i].weekday() == dates[day_index].weekday()][-4:]
        source = np.mean(history, axis=0) if history else np.mean(load[:day_index], axis=0)
    padded = np.pad(np.asarray(source, dtype=float), (1, 1), mode="edge")
    return np.convolve(padded, np.ones(3) / 3.0, mode="valid")


def pv_trend_forecast(day_index: int, pv: np.ndarray, typical_pv: np.ndarray) -> np.ndarray:
    """Causal linear trend over the last 21 completed PV days."""
    if day_index == 0:
        return np.maximum(typical_pv.copy(), 0.0)
    history = pv[max(0, day_index - 21):day_index]
    if len(history) == 1:
        return np.maximum(history[0].copy(), 0.0)
    x = np.arange(len(history), dtype=float)
    centered = x - x.mean()
    slope = centered @ np.asarray(history, dtype=float) / (centered @ centered)
    return np.maximum(np.mean(history, axis=0) + slope * (len(history) - x.mean()), 0.0)


def asof_pv_forecast(day_index: int, issue_hour: int, dates, pv, forecasts) -> np.ndarray:
    """Interpolate the forecast available at issue_hour without crossing the day."""
    if issue_hour == 0:
        anchor = float(pv[day_index - 1, -1]) if day_index > 0 else 0.0
    else:
        anchor = float(pv[day_index, issue_hour * 6 - 1])
    values = np.asarray(forecasts[dates[day_index]][issue_hour], dtype=float)
    knots = issue_hour + np.arange(25, dtype=float)
    endpoint_hours = (np.arange(T, dtype=float) + 1.0) / 6.0
    result = np.interp(endpoint_hours, knots, np.r_[anchor, values])
    result[endpoint_hours <= issue_hour] = 0.0
    return np.maximum(result, 0.0)


def q3_raw_net(day_index: int, issue_hour: int, dates, load, pv, forecasts, typical_load, typical_pv, blend: float) -> np.ndarray:
    load_f = load_history_forecast(day_index, dates, load, typical_load)
    pv_hist = pv_trend_forecast(day_index, pv, typical_pv)
    pv_issue = asof_pv_forecast(day_index, issue_hour, dates, pv, forecasts)
    pv_f = blend * pv_issue + (1.0 - blend) * pv_hist
    return (load_f - pv_f) * DT


def q3_design_net(day_index: int, issue_hour: int, dates, load, pv, forecasts, typical_load, typical_pv, params: Q34Parameters) -> np.ndarray:
    raw = q3_raw_net(day_index, issue_hour, dates, load, pv, forecasts, typical_load, typical_pv, params.blend)
    if day_index == 0:
        return raw
    residuals = []
    start = max(1, day_index - params.load_error_window)
    for j in range(start, day_index):
        residuals.append((load[j] - pv[j]) * DT - q3_raw_net(j, issue_hour, dates, load, pv, forecasts, typical_load, typical_pv, params.blend))
    if not residuals:
        return raw
    return raw + np.quantile(np.asarray(residuals), params.alpha, axis=0, method="linear")


def _make_lp(
    price: np.ndarray,
    load_kwh: np.ndarray,
    pv_kwh: np.ndarray,
    initial_soc: float,
    terminal_value: float,
    base_purchase: np.ndarray | None = None,
    terminal_soc_target: float | None = None,
    allowed_addition_hours: tuple[int, ...] | None = None,
    addition_start_slot: int = 0,
    no_adjustment: bool = False,
) -> dict[str, np.ndarray]:
    """Solve either the day-ahead LP or the upward-only revision LP."""
    n_t = len(price)
    net = np.asarray(load_kwh, dtype=float) - np.asarray(pv_kwh, dtype=float)
    if len(net) != n_t:
        raise ValueError("LP input length mismatch")
    use_adjustment = base_purchase is not None

    if not use_adjustment:
        # Reference day-ahead layout: [q, charge, discharge, soc, surplus].
        q, ch, dis, soc, surplus = 0, n_t, 2 * n_t, 3 * n_t, 4 * n_t
        n = 5 * n_t
        objective = np.zeros(n)
        objective[q:q+n_t] = price
        objective[ch:ch+n_t] = REGULARIZER
        objective[dis:dis+n_t] = REGULARIZER
        objective[soc+n_t-1] -= terminal_value
        bounds = ([(0.0, None)] * n_t +
                  [(0.0, P_MAX * DT)] * (2 * n_t) +
                  [(SOC_MIN, SOC_MAX)] * n_t +
                  [(0.0, None)] * n_t)
    else:
        # Reference revision layout: [up, charge, discharge, emergency, soc, surplus].
        up, ch, dis, em, soc, surplus = (0, n_t, 2*n_t, 3*n_t, 4*n_t, 5*n_t)
        n = 6 * n_t
        objective = np.zeros(n)
        objective[up:up+n_t] = 1.5 * price
        objective[ch:ch+n_t] = REGULARIZER
        objective[dis:dis+n_t] = REGULARIZER
        objective[em:em+n_t] = 5.0 * price
        objective[soc+n_t-1] -= terminal_value
        if allowed_addition_hours is None:
            up_bounds = [(0.0, None)] * n_t
        else:
            up_bounds = [
                (0.0, None)
                if ((((addition_start_slot + t) // 36) * 6) in allowed_addition_hours)
                else (0.0, 0.0)
                for t in range(n_t)
            ]
        if no_adjustment:
            up_bounds = [(0.0, 0.0)] * n_t
        bounds = (up_bounds +
                  [(0.0, P_MAX * DT)] * (2 * n_t) +
                  [(0.0, None)] * n_t +
                  [(SOC_MIN, SOC_MAX)] * n_t +
                  [(0.0, None)] * n_t)

    rows: list[int] = []
    cols: list[int] = []
    vals: list[float] = []
    rhs: list[float] = []

    def add(row: int, col: int, value: float) -> None:
        rows.append(row)
        cols.append(col)
        vals.append(value)

    row = 0
    for t in range(n_t):
        if use_adjustment:
            add(row, up+t, 1.0)
            add(row, em+t, 1.0)
            rhs.append(float(net[t] - base_purchase[t]))
        else:
            add(row, q+t, 1.0)
            rhs.append(float(net[t]))
        add(row, dis+t, 1.0)
        add(row, ch+t, -1.0)
        add(row, surplus+t, -1.0)
        row += 1

    if terminal_soc_target is not None:
        add(row, soc + n_t - 1, 1.0)
        rhs.append(float(terminal_soc_target))
        row += 1
    for t in range(n_t):
        add(row, soc+t, 1.0)
        add(row, ch+t, -ETA)
        add(row, dis+t, 1.0 / ETA)
        if t:
            add(row, soc+t-1, -1.0)
            rhs.append(0.0)
        else:
            rhs.append(float(initial_soc))
        row += 1

    result = linprog(
        objective,
        A_eq=coo_matrix((vals, (rows, cols)), shape=(row, n)).tocsr(),
        b_eq=np.asarray(rhs, dtype=float),
        bounds=bounds,
        method="highs",
    )
    if not result.success:
        raise RuntimeError(f"滚动 LP 求解失败: {result.message}")
    x = result.x
    return {
        "purchase_kwh": (np.maximum(x[q:q+n_t], 0.0) if not use_adjustment
                         else np.asarray(base_purchase, dtype=float) + np.maximum(x[up:up+n_t], 0.0)),
        "emergency_kwh": np.maximum(x[em:em+n_t], 0.0) if use_adjustment else np.zeros(n_t),
        "charge_kwh": np.maximum(x[ch:ch+n_t], 0.0),
        "discharge_kwh": np.maximum(x[dis:dis+n_t], 0.0),
        "soc": x[soc:soc+n_t],
        "surplus_kwh": np.maximum(x[surplus:surplus+n_t], 0.0),
        "up_kwh": np.maximum(x[up:up+n_t], 0.0) if use_adjustment else np.zeros(n_t),
        "down_kwh": np.zeros(n_t),
        "objective": float(result.fun),
    }


def _execute_slot(
    state: float,
    actual_load_kwh: float,
    actual_pv_kwh: float,
    purchase_kwh: float,
    reference_soc: float,
    reserve_ratio: float,
) -> tuple[float, float, float, float, float, float]:
    balance = float(purchase_kwh) - float(actual_load_kwh - actual_pv_kwh)
    if balance >= 0.0:
        charge = min(balance, P_MAX * DT, max((SOC_MAX - state) / ETA, 0.0))
        surplus = balance - charge
        discharge = 0.0
        emergency = 0.0
    else:
        charge = 0.0
        reserve = SOC_MIN + reserve_ratio * max(reference_soc - SOC_MIN, 0.0)
        discharge = min(-balance, P_MAX * DT, max(ETA * (state - reserve), 0.0))
        surplus = 0.0
        emergency = -balance - discharge
    state = float(np.clip(state + ETA * charge - discharge / ETA, SOC_MIN, SOC_MAX))
    return state, charge, discharge, emergency, surplus, balance


def _load_forecast(day_index: int, dates, load, pv, typical_load, typical_pv, params) -> np.ndarray:
    load_f, _, _ = forecast_day(day_index, dates, load, pv, typical_load, typical_pv, Q2Parameters(error_window=params.error_window))
    return np.maximum(load_f, 0.0)


def forecast_price(
    day_index: int,
    dates,
    price_matrix: dict[date, np.ndarray],
    typical_price: np.ndarray,
    window: int = 4,
    method: str = "weighted",
) -> np.ndarray:
    """Causal day-ahead price forecast from up to four prior same-weekdays."""
    if day_index == 0:
        return typical_price.copy()
    weekday = dates[day_index].weekday()
    history = [price_matrix[dates[i]] for i in range(day_index) if dates[i].weekday() == weekday][-window:]
    if not history:
        history = [typical_price]
    values = np.asarray(history)
    if method == "weighted":
        weights = np.arange(1, len(history) + 1, dtype=float)
        weights /= weights.sum()
        return np.average(values, axis=0, weights=weights)
    if method == "mean":
        return np.mean(values, axis=0)
    if method == "median":
        return np.median(values, axis=0)
    raise ValueError(f"unknown price forecast method: {method}")


def simulate_q3_or_q4(variable_price: bool, params: Q34Parameters | None = None) -> tuple[dict, dict]:
    dates, load, pv = load_attachment2()
    att1 = load_attachment1(DATA / "附件1.xlsx")
    prices = np.asarray(att1["price"], dtype=float)
    if variable_price:
        wb = load_workbook(DATA / "附件4.xlsx", read_only=True, data_only=True)
        ws = wb.active
        price_by_date = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            price_by_date[_as_date(row[0])] = np.asarray([float(x or 0.0) for x in row[1:T+1]], dtype=float)
        wb.close()
    else:
        price_by_date = {d: prices for d in dates}
    forecasts = load_attachment3()
    typical_load = np.asarray(att1["load"], dtype=float)
    typical_pv = np.asarray(att1["pv"], dtype=float)
    typical_price = np.asarray(att1["price"], dtype=float)
    params = params or Q34Parameters()
    warmup_params = WARMUP_PARAMETERS
    state = INITIAL_SOC
    records = []

    for i, day in enumerate(dates):
        active_q2 = warmup_params if day < REPORT_START else Q2Parameters(
            alpha=params.alpha,
            error_window=params.load_error_window,
            reserve_ratio=params.reserve_ratio,
            terminal_value=params.terminal_value,
        )
        active_terminal = active_q2.terminal_value
        active_reserve = active_q2.reserve_ratio
        actual_price = price_by_date[day]
        plan_price = forecast_price(i, dates, price_by_date, prices, params.price_window, params.price_method)
        price = actual_price
        if day < REPORT_START:
            load_f, pv_f, residuals = forecast_day(
                i, dates, load, pv, typical_load, typical_pv, active_q2,
                forecast_variant="official",
            )
            design = design_demand(load_f, pv_f, residuals, active_q2.alpha)
            warmup_plan = plan_day(plan_price, design, state, warmup_params)
            q0_plan = {"purchase_kwh": warmup_plan["purchase_kw"], "soc": warmup_plan["soc"]}
        else:
            design = q3_design_net(i, 0, dates, load, pv, forecasts, typical_load, typical_pv, params)
            q0_plan = _make_lp(
                plan_price, design, np.zeros(T), state, active_terminal,
                terminal_soc_target=params.terminal_soc_target,
            )
        q0 = q0_plan["purchase_kwh"]
        final_q = q0.copy()
        state_day = float(state)
        ch = np.zeros(T); dis = np.zeros(T); em = np.zeros(T); surplus = np.zeros(T); soc = np.zeros(T)
        reference_soc = q0_plan["soc"].copy()
        adj_up = np.zeros(T); adj_down = np.zeros(T)
        decisions = []
        for t in range(T):
            if day >= REPORT_START and t in params.adjustment_times:
                issue_hour = t // 6
                update_design = q3_design_net(i, issue_hour, dates, load, pv, forecasts, typical_load, typical_pv, params)
                expected_price = plan_price[t:].copy()
                adjustment = _make_lp(
                    expected_price, update_design[t:], np.zeros(T - t), state_day, active_terminal,
                    base_purchase=q0[t:],
                    terminal_soc_target=params.terminal_soc_target,
                    allowed_addition_hours=tuple(x // 6 for x in params.adjustment_times),
                    addition_start_slot=t,
                )
                fallback = _make_lp(
                    expected_price,
                    update_design[t:],
                    np.zeros(T - t),
                    state_day,
                    active_terminal,
                    base_purchase=q0[t:],
                    terminal_soc_target=params.terminal_soc_target,
                    allowed_addition_hours=tuple(x // 6 for x in params.adjustment_times),
                    addition_start_slot=t,
                    no_adjustment=True,
                )
                end = min(t + 36, T)
                proxy_inconsistency = adjustment["objective"] > fallback["objective"] + 1e-5
                accepted = (
                    not proxy_inconsistency
                    and adjustment["objective"] < fallback["objective"] - 1e-5
                    and float(adjustment["up_kwh"][:36].sum()) > 1e-7
                )
                if accepted:
                    final_q[t:end] = q0[t:end] + adjustment["up_kwh"][: end - t]
                    reference_soc[t:end] = adjustment["soc"][: end - t]
                    adj_up[t:end] = adjustment["up_kwh"][: end - t]
                decisions.append({
                    "issue_hour": issue_hour, "initial_soc": float(state_day),
                    "proxy_adjust_cost": adjustment["objective"],
                    "proxy_no_adjust_cost": fallback["objective"],
                    "accepted": bool(accepted),
                    "committed_addition": float(adj_up[t:end].sum()),
                    "planned_end_soc": float(adjustment["soc"][-1]),
                    "proxy_inconsistency": bool(proxy_inconsistency),
                })
            state_day, ch[t], dis[t], em[t], surplus[t], _ = _execute_slot(
                state_day,
                load[i, t] * DT,
                pv[i, t] * DT,
                final_q[t],
                reference_soc[t],
                active_reserve,
            )
            soc[t] = state_day
        q0_cost = float(np.dot(actual_price, q0))
        adjustment_cost = float(np.dot(1.5 * price, adj_up) + np.dot(0.5 * price, adj_down))
        emergency_cost = float(np.dot(5.0 * actual_price, em))
        records.append({
            "date": day.isoformat(),
            "soc0": float(state), "soc24": float(state_day),
            "price": actual_price.tolist(), "plan_price": plan_price.tolist(), "purchase_kwh": q0.tolist(), "adjusted_purchase_kwh": final_q.tolist(),
            "adjustment_up_kwh": adj_up.tolist(), "adjustment_down_kwh": adj_down.tolist(),
            "charge_kwh": ch.tolist(), "discharge_kwh": dis.tolist(), "emergency_kwh": em.tolist(),
            "surplus_kwh": surplus.tolist(), "soc_trace": soc.tolist(),
            "reference_soc": reference_soc.tolist(),
            "day_ahead_reference_soc": q0_plan["soc"].tolist(),
            "adjustment_decisions": decisions,
            "load_kwh": (load[i] * DT).tolist(), "pv_kwh": (pv[i] * DT).tolist(),
            "plan_cost": q0_cost, "adjustment_cost": adjustment_cost, "emergency_cost": emergency_cost,
            "total_cost": q0_cost + adjustment_cost + emergency_cost,
        })
        state = state_day

    report = [r for r in records if r["date"] >= REPORT_START.isoformat()]
    summary = {
        "variable_price": variable_price,
        "days": len(report), "intervals": len(report) * T,
        "plan_kwh": sum(sum(r["purchase_kwh"]) for r in report),
        "adjusted_plan_kwh": sum(sum(r["adjusted_purchase_kwh"]) for r in report),
        "adjustment_up_kwh": sum(sum(r["adjustment_up_kwh"]) for r in report),
        "adjustment_down_kwh": sum(sum(r["adjustment_down_kwh"]) for r in report),
        "emergency_kwh": sum(sum(r["emergency_kwh"]) for r in report),
        "plan_cost": sum(r["plan_cost"] for r in report),
        "adjustment_cost": sum(r["adjustment_cost"] for r in report),
        "emergency_cost": sum(r["emergency_cost"] for r in report),
        "total_cost": sum(r["total_cost"] for r in report),
        "january_end_soc": report[0]["soc0"], "year_end_soc": report[-1]["soc24"],
    }
    return {"records": report, "prices": prices.tolist()}, summary


def simulate_q4_2(
    price_window: int = 4,
    price_method: str = "weighted",
) -> tuple[dict, dict]:
    """Question 2 policy with the daily volatile prices from attachment 4."""
    dates, load, pv = load_attachment2()
    att1 = load_attachment1(DATA / "附件1.xlsx")
    typical_load = np.asarray(att1["load"], dtype=float)
    typical_pv = np.asarray(att1["pv"], dtype=float)
    typical_price = np.asarray(att1["price"], dtype=float)
    price_by_date: dict[date, np.ndarray] = {}
    wb = load_workbook(DATA / "附件4.xlsx", read_only=True, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            price_by_date[_as_date(row[0])] = np.asarray([float(x or 0.0) for x in row[1:T+1]], dtype=float)
    wb.close()
    params = Q2Parameters()
    warmup_params = WARMUP_PARAMETERS
    state = INITIAL_SOC
    all_records = []
    for i, day in enumerate(dates):
        active_params = warmup_params if day < REPORT_START else params
        actual_price = price_by_date[day]
        price = forecast_price(i, dates, price_by_date, typical_price, price_window, price_method)
        load_f, pv_f, residuals = forecast_day(
            i, dates, load, pv, typical_load, typical_pv, active_params,
            forecast_variant="official" if day < REPORT_START else "optimized",
        )
        demand = design_demand(load_f, pv_f, residuals, active_params.alpha)
        plan = plan_day(
            price,
            demand,
            state,
            active_params,
            terminal_soc_target=None if day < REPORT_START else TERMINAL_SOC_TARGET,
        )
        execution = execute_day(actual_price, load[i], pv[i], plan["purchase_kw"], plan["soc"], state, active_params.reserve_ratio)
        all_records.append({
            "date": day.isoformat(), "soc0": float(state), "soc24": float(execution["soc_trace"][-1]),
            "price": actual_price.tolist(), "plan_price": price.tolist(), "purchase_kwh": plan["purchase_kw"].tolist(),
            "charge_kwh": execution["charge_kwh"].tolist(), "discharge_kwh": execution["discharge_kwh"].tolist(),
            "emergency_kwh": execution["emergency_kwh"].tolist(), "surplus_kwh": execution["surplus_kwh"].tolist(),
            "soc_trace": execution["soc_trace"].tolist(), "load_kwh": (load[i] * DT).tolist(),
            "reference_soc": plan["soc"].tolist(),
            "pv_kwh": (pv[i] * DT).tolist(),
            "plan_cost": float(np.dot(actual_price, plan["purchase_kw"])),
            "emergency_cost": float(np.dot(5.0 * actual_price, execution["emergency_kwh"])),
        })
        all_records[-1]["total_cost"] = all_records[-1]["plan_cost"] + all_records[-1]["emergency_cost"]
        state = float(execution["soc_trace"][-1])
    report = [r for r in all_records if r["date"] >= REPORT_START.isoformat()]
    summary = {
        "variable_price": True, "days": len(report), "intervals": len(report) * T,
        "plan_kwh": sum(sum(r["purchase_kwh"]) for r in report),
        "emergency_kwh": sum(sum(r["emergency_kwh"]) for r in report),
        "plan_cost": sum(r["plan_cost"] for r in report),
        "emergency_cost": sum(r["emergency_cost"] for r in report),
        "total_cost": sum(r["total_cost"] for r in report),
        "january_end_soc": report[0]["soc0"], "year_end_soc": report[-1]["soc24"],
    }
    return {"records": report}, summary


def validate_q4_2(bundle: dict, summary: dict) -> dict:
    max_balance = 0.0; max_soc_error = 0.0; min_soc = float("inf"); max_soc = float("-inf")
    for i, r in enumerate(bundle["records"]):
        q = np.asarray(r["purchase_kwh"]); ch = np.asarray(r["charge_kwh"]); dis = np.asarray(r["discharge_kwh"]); em = np.asarray(r["emergency_kwh"]); soc = np.asarray(r["soc_trace"])
        balance = q + np.asarray(r["pv_kwh"]) + dis + em - np.asarray(r["load_kwh"]) - ch - np.asarray(r["surplus_kwh"])
        expected = np.empty(T); expected[0] = r["soc0"] + ETA * ch[0] - dis[0] / ETA; expected[1:] = soc[:-1] + ETA * ch[1:] - dis[1:] / ETA
        max_balance = max(max_balance, float(np.max(np.abs(balance)))); max_soc_error = max(max_soc_error, float(np.max(np.abs(soc - expected))))
        min_soc = min(min_soc, float(soc.min()), r["soc0"]); max_soc = max(max_soc, float(soc.max()), r["soc0"])
        if i and abs(r["soc0"] - bundle["records"][i-1]["soc24"]) > 1e-8: raise AssertionError("cross-day SOC mismatch")
        if soc.min() < SOC_MIN - 1e-8 or soc.max() > SOC_MAX + 1e-8: raise AssertionError("SOC bound violation")
        if np.max(ch * dis) > 1e-8: raise AssertionError("simultaneous charge/discharge")
    if max_balance > 1e-8 or max_soc_error > 1e-8: raise AssertionError("balance/SOC mismatch")
    return {"status": "PASS", "intervals": len(bundle["records"]) * T, "max_balance_error_kwh": max_balance, "max_soc_error_kwh": max_soc_error, "soc_range_kwh": [min_soc, max_soc]}


def _segments(values: np.ndarray) -> list[tuple[str, float]]:
    out = []
    start = None
    for t in range(T + 1):
        active = t < T and values[t] > EMERGENCY_EPS
        if active and start is None:
            start = t
        if not active and start is not None:
            out.append((interval_label(start, t), float(values[start:t].sum())))
            start = None
    return out


def validate(bundle: dict, summary: dict) -> dict:
    records = bundle["records"]
    max_balance = 0.0; max_soc_error = 0.0; min_soc = float("inf"); max_soc = float("-inf")
    for i, r in enumerate(records):
        if i and abs(r["soc0"] - records[i-1]["soc24"]) > 1e-8: raise AssertionError("cross-day SOC mismatch")
        q = np.asarray(r["adjusted_purchase_kwh"]); ch = np.asarray(r["charge_kwh"]); dis = np.asarray(r["discharge_kwh"]); em = np.asarray(r["emergency_kwh"]); soc = np.asarray(r["soc_trace"])
        balance = q + np.asarray(r["pv_kwh"]) + dis + em - np.asarray(r["load_kwh"]) - ch - np.asarray(r["surplus_kwh"])
        expected = np.empty(T); expected[0] = r["soc0"] + ETA * ch[0] - dis[0] / ETA; expected[1:] = soc[:-1] + ETA * ch[1:] - dis[1:] / ETA
        max_balance = max(max_balance, float(np.max(np.abs(balance)))); max_soc_error = max(max_soc_error, float(np.max(np.abs(soc - expected))))
        min_soc = min(min_soc, float(soc.min()), r["soc0"]); max_soc = max(max_soc, float(soc.max()), r["soc0"])
        if np.min(np.r_[q, ch, dis, em]) < -1e-8 or np.max(ch) > P_MAX * DT + 1e-8 or np.max(dis) > P_MAX * DT + 1e-8: raise AssertionError("flow bound violation")
        if soc.min() < SOC_MIN - 1e-8 or soc.max() > SOC_MAX + 1e-8: raise AssertionError("SOC bound violation")
        if np.max(ch * dis) > 1e-8: raise AssertionError("simultaneous charge/discharge")
        if abs(r["total_cost"] - r["plan_cost"] - r["adjustment_cost"] - r["emergency_cost"]) > 1e-7: raise AssertionError("cost mismatch")
    if max_balance > 1e-8 or max_soc_error > 1e-8: raise AssertionError("balance/SOC mismatch")
    return {"status": "PASS", "intervals": len(records) * T, "max_balance_error_kwh": max_balance, "max_soc_error_kwh": max_soc_error, "soc_range_kwh": [min_soc, max_soc]}


def save_q34_workbook(
    bundle: dict,
    output_name: str,
    template_name: str,
    adjusted: bool,
    plan_sheet_total_cost: bool = False,
) -> Path:
    """Fill one supplied result template from the freshly simulated records."""
    template = DATA / "附件5" / template_name
    wb = load_workbook(template)
    records = bundle["records"]
    by_date = {date.fromisoformat(r["date"]): r for r in records}
    windows = [
        (list(range(0, 24)), "0:00-4:00"),
        (list(range(24, 48)), "4:00-8:00"),
        (list(range(48, 72)), "8:00-12:00"),
        (list(range(72, 96)), "12:00-16:00"),
        (list(range(96, 120)), "16:00-20:00"),
        (list(range(120, 144)), "20:00-24:00"),
    ]

    ws_plan = wb["计划购电量"]
    for name in (["计划购电量", "调整购电量"] if adjusted else ["计划购电量"]):
        for t in range(T):
            wb[name].cell(1, 2 + t, interval_label(t, t + 1))
    rows = {}
    for row in range(2, ws_plan.max_row + 1):
        value = ws_plan.cell(row, 1).value
        if isinstance(value, datetime):
            rows[value.date()] = row
        elif isinstance(value, date):
            rows[value] = row
    for day, record in by_date.items():
        row = rows[day]
        values = record["purchase_kwh"]
        for t, value in enumerate(values):
            ws_plan.cell(row, 2 + t, round(float(value), 6))
        ws_plan.cell(row, 146, round(float(sum(values)), 6))
        plan_sheet_cost = float(record["total_cost"] if plan_sheet_total_cost else record["plan_cost"])
        ws_plan.cell(row, 147, round(plan_sheet_cost, 6))

    if adjusted:
        ws_adjusted = wb["调整购电量"]
        for day, record in by_date.items():
            row = rows[day]
            for t, value in enumerate(record["adjusted_purchase_kwh"]):
                ws_adjusted.cell(row, 2 + t, round(float(value), 6))
            ws_adjusted.cell(row, 146, round(float(sum(record["adjusted_purchase_kwh"])), 6))
            ws_adjusted.cell(row, 147, round(float(record["total_cost"]), 6))

    ws_storage = wb["充放电量"]
    ws_storage.delete_rows(2, ws_storage.max_row)
    row = 2
    for day, record in by_date.items():
        charge = record["charge_kwh"]
        discharge = record["discharge_kwh"]
        for block, (indices, label) in enumerate(windows):
            ws_storage.cell(row, 1, datetime.combine(day, datetime.min.time()) if block == 0 else None)
            ws_storage.cell(row, 2, label)
            ws_storage.cell(row, 3, round(float(sum(charge[i] for i in indices)), 6))
            ws_storage.cell(row, 4, round(float(sum(discharge[i] for i in indices)), 6))
            if block == 0:
                ws_storage.cell(row, 5, "0:00")
                ws_storage.cell(row, 6, round(float(record["soc0"]), 6))
            elif block == 1:
                ws_storage.cell(row, 5, "24:00")
                ws_storage.cell(row, 6, round(float(record["soc24"]), 6))
            row += 1

    ws_emergency = wb["紧急购电量"]
    ws_emergency.delete_rows(2, ws_emergency.max_row)
    row = 2
    for day, record in by_date.items():
        events = _segments(np.asarray(record["emergency_kwh"], dtype=float))
        if not events:
            events = [("无", 0.0)]
        for index, (span, quantity) in enumerate(events):
            ws_emergency.cell(row, 1, datetime.combine(day, datetime.min.time()) if index == 0 else None)
            ws_emergency.cell(row, 2, span)
            ws_emergency.cell(row, 3, round(float(quantity), 6))
            row += 1

    for ws in wb.worksheets:
        ws.freeze_panes = "B2"
        for row_cells in ws.iter_rows(min_row=2):
            for cell in row_cells:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000000"
                elif isinstance(cell.value, datetime):
                    cell.number_format = "yyyy/mm/dd"
    output = ROOT / output_name
    wb.save(output)
    return output


def save_json(bundle: dict, summary: dict, validation: dict, name: str, parameters: dict | None = None) -> None:
    payload = {"parameters": parameters or asdict(Q34Parameters()),
               "warmup_parameters": asdict(WARMUP_PARAMETERS),
               "selection_scope": "Frozen after retrospective comparisons; causal decisions, not independent held-out selection",
               "summary": summary, "validation": validation, **bundle}
    (ROOT / name).write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("--no-xlsx", action="store_true")
    args = parser.parse_args()
    for variable_price, data_name, summary_name in [(False, "q3_rebuild_data.json", "q3_summary.json"), (True, "q4_3_rebuild_data.json", "q4_3_summary.json")]:
        bundle, summary = simulate_q3_or_q4(variable_price)
        validation = validate(bundle, summary)
        save_json(bundle, summary, validation, data_name, asdict(Q34Parameters()))
        if not args.no_xlsx:
            save_q34_workbook(bundle, "result3.xlsx" if not variable_price else "result4-3.xlsx", "result3.xlsx" if not variable_price else "result4-3.xlsx", adjusted=True)
        (ROOT / summary_name).write_text(json.dumps({"parameters": asdict(Q34Parameters()), "summary": summary, "validation": validation}, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({"variable_price": variable_price, "summary": summary, "validation": validation}, ensure_ascii=False))
    bundle, summary = simulate_q4_2()
    validation = validate_q4_2(bundle, summary)
    q4_2_parameters = asdict(Q2Parameters())
    save_json(bundle, summary, validation, "q4_2_rebuild_data.json", q4_2_parameters)
    if not args.no_xlsx:
        save_q34_workbook(bundle, "result4-2.xlsx", "result4-2.xlsx", adjusted=False, plan_sheet_total_cost=True)
    (ROOT / "q4_2_summary.json").write_text(json.dumps({"parameters": q4_2_parameters, "summary": summary, "validation": validation}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"q4_2": summary, "validation": validation}, ensure_ascii=False))


if __name__ == "__main__":
    main()
